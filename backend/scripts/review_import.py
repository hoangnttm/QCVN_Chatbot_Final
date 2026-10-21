# ruff: noqa: E501
"""Import an expert-filled review pack (Excel) into PostgreSQL (PLAN.md §11).

    python backend/scripts/review_import.py data/review/review-pack-v1-<date>.xlsx --reviewer "Tên" [--approve-tables] [--dry-run]

For every row whose KẾT QUẢ column is filled, a `review_decisions` record is written and the
object's `review_status` is updated through `app.review.service` (the same code path as the admin
review UI). With `--approve-tables`, a threshold table whose sampled rows are all "Đúng" gets its
remaining `pending` rules bulk-approved (the 10 % sample rule is checked by the service).
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core.config import Settings  # noqa: E402
from app.db.session import get_sessionmaker  # noqa: E402
from app.review import service  # noqa: E402

DECISION = {"Đúng": "approved", "Sai": "rejected", "Cần làm rõ": "needs_recheck"}
THRESHOLD_SHEETS = {
    "2a-QCVN40 B1": "QCVN40_2025.B1",
    "2b-QCVN40 B2": "QCVN40_2025.B2",
    "2c-QCVN19 B1": "QCVN19_2024.B1",
    "2d-QCVN19 B2": "QCVN19_2024.B2",
}
TARGET_BY_PREFIX = [
    ("rel:", "document_relations"),
    ("fact:", "registry_facts"),
    ("map:", "dictionary_mapping"),
    ("TT10.MR.", "monitoring_requirements"),
]


def _target_type(sheet: str, key: str) -> str:
    for prefix, t in TARGET_BY_PREFIX:
        if key.startswith(prefix):
            return t
    if sheet == "7-Chuyển tiếp":
        return "transition_rules"
    return "table_rows"


def read_decisions(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True)
    out: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        if ws.title.startswith(("0-", "1-")):
            continue
        header_row = None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), start=1):
            if row and row[0] in ("Mã dòng", "Mã", "Văn bản / chủ đề"):
                header_row = i
                headers = list(row)
                break
        if header_row is None or "KẾT QUẢ" not in headers:
            continue
        c_res = headers.index("KẾT QUẢ")
        c_fix = headers.index("Sửa thành / giá trị đúng")
        c_note = headers.index("Ghi chú chuyên gia")
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not row or row[0] is None:
                continue
            result = (row[c_res] or "").strip() if isinstance(row[c_res], str) else row[c_res]
            out.append(
                {
                    "sheet": ws.title,
                    "key": str(row[0]),
                    "result": result or None,
                    "decision": DECISION.get(result or ""),
                    "correction": row[c_fix],
                    "note": row[c_note],
                }
            )
    return out


def apply(
    decisions: list[dict[str, Any]],
    reviewer: str,
    source: str,
    approve_tables: bool,
    settings: Settings,
) -> dict[str, Any]:
    maker = get_sessionmaker(settings)
    filled = [d for d in decisions if d["decision"]]
    summary: dict[str, Any] = {"rows_read": len(decisions), "decisions": len(filled), "bulk": []}
    with maker.begin() as session:
        for d in filled:
            service.apply_decision(
                session,
                target_type=_target_type(d["sheet"], d["key"]),
                target_id=d["key"],
                decision=d["decision"],
                reviewer=reviewer,
                note=str(d["note"]) if d["note"] else None,
                correction=str(d["correction"]) if d["correction"] else None,
                source=f"{source}#{d['sheet']}",
                check_exists=False,  # sheets may reference rows of a document not loaded here
            )
        if approve_tables:
            session.flush()
            for sheet, table_code in THRESHOLD_SHEETS.items():
                sample = [d for d in decisions if d["sheet"] == sheet]
                if not sample or any(d["decision"] != "approved" for d in sample):
                    continue
                try:
                    res = service.bulk_approve(
                        session, table_code, reviewer, source=f"{source}#{sheet}"
                    )
                except service.BulkApproveBlocked:
                    continue
                summary["bulk"].append(
                    {"table": table_code, "sample": len(sample), "rules": res["rules"]}
                )
        status = service.rules_by_status(session)
    summary["limit_rules_by_status"] = status
    return summary


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--reviewer", required=True)
    ap.add_argument("--approve-tables", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    path = Path(args.workbook)
    decisions = read_decisions(path)
    by_sheet: dict[str, dict[str, int]] = {}
    for d in decisions:
        s = by_sheet.setdefault(d["sheet"], {"Đúng": 0, "Sai": 0, "Cần làm rõ": 0, "trống": 0})
        s[d["result"] if d["result"] in s else "trống"] += 1
    for sheet, s in by_sheet.items():
        print(f"{sheet:<22} " + "  ".join(f"{k}={v}" for k, v in s.items()))
    filled = sum(1 for d in decisions if d["decision"])
    if not filled:
        print("Không có ô KẾT QUẢ nào được điền.")
        return 1
    if args.dry_run:
        print(f"dry-run: {filled} quyết định sẽ được ghi.")
        return 0
    settings = Settings()
    summary = apply(decisions, args.reviewer, path.name, args.approve_tables, settings)
    log = settings.data_dir / "review" / f"decisions-{dt.date.today().isoformat()}.json"
    log.write_text(
        json.dumps(
            {
                "workbook": path.name,
                "reviewer": args.reviewer,
                "summary": summary,
                "decisions": decisions,
            },
            ensure_ascii=False,
            indent=1,
            default=str,
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=1))
    print(f"→ {log}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
