"""Build the expert review pack (Excel) from the loaded PostgreSQL data — PLAN.md §11, Phase 1 gate
"chuyên gia kiểm 30 dòng/bảng ngưỡng, sai số cho phép = 0".

Usage:
    python backend/scripts/review_pack.py [--out data/review/review-pack-v1.xlsx] [--sample 30]

Sheets (Vietnamese, reviewer fills the yellow columns only):
    0-Hướng dẫn, 1-Tổng hợp, 2a..2d threshold tables (sampled rows, values per zone),
    3-Phương pháp, 3b-TT10 kỹ thuật, 4-Ngành PL2, 5-Hệ số TEF, 6-Quan hệ, 7-Chuyển tiếp,
    8-Sự kiện đăng ký, 9-Ánh xạ từ điển, 10-Quan trắc TT10, 12-Kịch bản áp dụng (engine chạy live),
    11-Phát hiện
Sampling is deterministic (risk-first, then evenly spaced) so re-runs give the same rows.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.worksheet import Worksheet  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.core.config import Settings  # noqa: E402
from app.db import models as m  # noqa: E402
from app.db.session import get_sessionmaker  # noqa: E402
from app.review.excel import (  # noqa: E402
    REVIEW_COLS,
    add_review_columns,
    finish_sheet,
    style_header,
)
from app.review.sampling import natural_key, sample_rows  # noqa: E402

THRESHOLD_TABLES = [
    ("2a-QCVN40 B1", "QCVN40_2025.B1"),
    ("2b-QCVN40 B2", "QCVN40_2025.B2"),
    ("2c-QCVN19 B1", "QCVN19_2024.B1"),
    ("2d-QCVN19 B2", "QCVN19_2024.B2"),
]
METHOD_TABLES = [
    "QCVN40_2025.PL1",
    "QCVN19_2024.PL1",
    "TT10_2021.B25",
    "TT10_2021.B26",
    "TT10_2021.B27",
    "TT10_2021.B28",
    "TT10_2021.B29",
    "TT10_2021.B30",
]
TT10_TECH_TABLES = ["TT10_2021.B3", "TT10_2021.B4", "TT10_2021.B5", "TT10_2021.B32"]
TEF_TABLES = ["QCVN19_2024.PL2", "QCVN19_2024.PL3"]

KNOWN_ISSUES = [
    (
        "TT_10_2021",
        "Phụ lục 5, Mục III, Khoản 3, Bảng 31 (khoảng cách điểm hút mẫu, % đường kính) và bảng ma "
        "trận ống khói chữ nhật chưa tách được thành dòng bảng (docling gộp thành đoạn văn). Hiện "
        "chỉ tra cứu được dưới dạng đoạn văn có trang.",
        "Chuyên gia xác nhận có cần tra cứu bảng này ở dạng số (theo tỉ lệ A/D, B/D) hay không.",
    ),
    (
        "TT_10_2021",
        "Phụ lục 6, Khoản 5: công thức quy đổi về nồng độ O₂ tham chiếu (Cstd) và đồ thị Hình 2/3 "
        "(số điểm hút mẫu) là hình ảnh trong PDF, chưa trích được thành text/số.",
        "Chuyên gia cung cấp công thức chuẩn (Cstd = Cm × (21 − %O₂std)/(21 − %O₂m)?) để nhập tay "
        "như registry fact có audit, hoặc chấp nhận chỉ dẫn tới trang PDF.",
    ),
    (
        "TT_10_2021",
        "Chương VII (quan trắc tự động, liên tục — CEMS) mới có bảng kỹ thuật B3, B4, B5 (sheet 3b); "
        "chưa tách MonitoringRequirement loại `cems` từ Đ33–38.",
        "Xác nhận có cần yêu cầu CEMS trong phạm vi MVP hay để Phase 3.",
    ),
    (
        "QCVN_19_2024",
        "Text QCVN 19:2024 lấy từ file Markdown chưa đối chiếu với bản ký (text_status = "
        "unverified); QCVN 19 đính kèm TT 45/2024 là bản dự thảo.",
        "PO cung cấp bản ký chính thức để đối chiếu (D13).",
    ),
    (
        "TT_45_2024 / TT_10_2021",
        "Hai văn bản là bản dự thảo: số hiệu, ngày ký, ngày hiệu lực do PO khai báo (registry fact) "
        "— xem sheet 8.",
        "Chuyên gia xác nhận các ngày khai báo; hệ thống luôn kèm cảnh báo 'bản dự thảo' khi trích dẫn.",
    ),
    (
        "Tần suất quan trắc",
        "TT10 Đ16.3/Đ17.5 chỉ dẫn 'theo quy định của pháp luật hiện hành' (NĐ 08/2022, NĐ 05/2025 "
        "không nạp — quyết định D10). Câu hỏi tần suất sẽ được từ chối kèm chỉ dẫn (M-01).",
        "Xác nhận cách trả lời này là chấp nhận được với người dùng.",
    ),
]


def cond_text(
    conds: list[dict[str, Any]], skip: tuple[str, ...] = ("zone", "equipment_path")
) -> str:
    out = []
    for c in conds:
        if c.get("type") in skip:
            continue
        label = c.get("text") or f"{c.get('type')} {c.get('op') or ''} {c.get('value')}".strip()
        if c.get("unit"):
            label += f" {c['unit']}"
        if c.get("inherited_from"):
            label += " (kế thừa từ dòng cha)"
        out.append(f"[{c.get('type')}] {label}")
    return "; ".join(out)


def value_text(v: dict[str, Any] | None) -> str:
    if not v:
        return ""
    if v.get("op") == "between":
        s = f"{v.get('min')} – {v.get('max')}"
    else:
        s = f"{v.get('op')} {v.get('value')}"
    if v.get("ref_o2") is not None:
        s += f" (O₂ {v['ref_o2']:g} %)"
    return s


SHEET_META: dict[str, tuple[str, int, int]] = {}  # sheet -> (result column letter, first, last)


def write_table_sheet(
    ws: Worksheet,
    title_lines: list[str],
    headers: list[str],
    rows: list[list[Any]],
    widths: list[int],
) -> tuple[int, int]:
    for line in title_lines:
        ws.append([line])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    ws.append([])
    ws.append(headers + REVIEW_COLS)
    header_row = ws.max_row
    ncols = len(headers) + len(REVIEW_COLS)
    style_header(ws, header_row, ncols)
    first = header_row + 1
    for r in rows:
        ws.append(r + ["", "", ""])
    last = ws.max_row
    add_review_columns(ws, header_row, first, last, ncols)
    finish_sheet(ws, widths + [14, 28, 36], header_row)
    SHEET_META[ws.title] = (get_column_letter(ncols - len(REVIEW_COLS) + 1), first, last)
    return first, last


# ----------------------------------------------------------------------------- data access


class Data:
    def __init__(self, session):
        self.s = session
        self.tables = {t.table_code: t for t in session.execute(select(m.TableDef)).scalars()}
        self.docs = {d.document_id: d for d in session.execute(select(m.Document)).scalars()}
        self._nodes: dict[str, m.StructureNode] = {}

    def rows(self, table_code: str) -> list[m.TableRow]:
        rows = (
            self.s.execute(select(m.TableRow).where(m.TableRow.table_code == table_code))
            .scalars()
            .all()
        )
        return sorted(rows, key=lambda r: (r.line_start, natural_key(r.row_id)))

    def rules_by_row(self, table_code: str) -> dict[str, list[m.LimitRule]]:
        out: dict[str, list[m.LimitRule]] = {}
        for r in self.s.execute(
            select(m.LimitRule).where(m.LimitRule.table_code == table_code)
        ).scalars():
            out.setdefault(r.row_id, []).append(r)
        return out

    def node(self, node_id: str) -> m.StructureNode | None:
        if node_id not in self._nodes:
            self._nodes[node_id] = self.s.get(m.StructureNode, node_id)
        return self._nodes[node_id]

    def footnotes(self, refs: list[str]) -> str:
        texts = []
        for ref in refs:
            n = self.node(ref)
            if n and n.text:
                texts.append(n.text)
        return " | ".join(texts)

    def source_ref(self, r) -> str:
        if r.printed_page:
            return f"tr. {r.printed_page} (PDF {r.pdf_page_start})"
        return f"dòng {r.line_start}" + (f"–{r.line_end}" if r.line_end != r.line_start else "")


# ----------------------------------------------------------------------------- sheets


def sheet_guide(wb: Workbook, today: str, sample: int) -> None:
    ws = wb.active
    ws.title = "0-Hướng dẫn"
    lines = [
        ("HỒ SƠ KIỂM TRA CHUYÊN GIA — Dữ liệu QCVN/Thông tư đã chuẩn hóa (v1)", True),
        (
            f"Sinh tự động ngày {today} từ cơ sở dữ liệu ingest (backend/scripts/review_pack.py).",
            False,
        ),
        ("", False),
        ("MỤC ĐÍCH", True),
        (
            "Xác nhận rằng các giá trị giới hạn, điều kiện áp dụng, phương pháp, quan hệ pháp lý và "
            "ngày hiệu lực mà hệ thống trích xuất tự động KHỚP 100 % với văn bản gốc. Chỉ dữ liệu "
            "được chuyên gia xác nhận mới được dùng để trả lời người dùng (PLAN.md §11: LimitRule, "
            "TransitionRule, REPLACES/REPEALS phải được duyệt).",
            False,
        ),
        ("", False),
        ("CÁCH KIỂM", True),
        ("1. Mở văn bản gốc (PDF/bản ký) theo cột 'Nhãn trích dẫn' và 'Vị trí nguồn'.", False),
        (
            "2. Đối chiếu từng cột hệ thống đã tách (thông số, đơn vị, thiết bị/điều kiện, giá trị theo "
            "cột A/B/C, O₂ tham chiếu, ngoại lệ, chú thích) với ô gốc trong văn bản.",
            False,
        ),
        (
            "3. Chỉ điền vào 3 cột màu vàng: KẾT QUẢ (chọn Đúng / Sai / Cần làm rõ), 'Sửa thành' "
            "(giá trị đúng theo văn bản nếu Sai) và 'Ghi chú'. Không sửa các cột khác.",
            False,
        ),
        (
            "4. Sai số cho phép = 0: bất kỳ sai lệch nào về số, đơn vị, cột áp dụng, điều kiện hay "
            "ngoại lệ đều ghi 'Sai'. Nghi ngờ cách hiểu (ví dụ điều kiện tách từ nhãn dòng) ghi "
            "'Cần làm rõ' và mô tả.",
            False,
        ),
        (
            f"5. Các bảng ngưỡng (sheet 2a–2d) là mẫu {sample} dòng/bảng chọn theo rủi ro (dòng có "
            "điều kiện, ngoại lệ, chú thích, kế thừa nhóm, khoảng giá trị) rồi rải đều; nếu mẫu "
            "đạt 0 sai, PO sẽ phê duyệt cả bảng (bulk approve, PLAN.md §11). Chuyên gia có thể "
            "yêu cầu thêm dòng.",
            False,
        ),
        (
            "6. Sheet 9 (ánh xạ từ điển) và 10 (yêu cầu quan trắc TT10) là xác nhận cách hiểu: "
            "tên trong văn bản có được gán đúng mã chuẩn/loại yêu cầu không.",
            False,
        ),
        (
            "7. Sheet 11 liệt kê điểm hệ thống chưa xử lý được — cần quyết định của chuyên gia/PO.",
            False,
        ),
        ("", False),
        ("Ý NGHĨA CỘT", True),
        ("Mã dòng / Mã: khóa nội bộ (giữ nguyên để nhập kết quả trở lại hệ thống).", False),
        ("Nhãn trích dẫn: nhãn hệ thống sẽ hiển thị cho người dùng — kiểm cả nhãn này.", False),
        (
            "Ô gốc: nội dung ô nguyên bản (Markdown/LaTeX) để đối chiếu nhanh khi không mở văn bản.",
            False,
        ),
        (
            "Giá trị: op giá trị, ví dụ '<= 50' hoặc '6 – 9' (khoảng); '(O₂ 6 %)' = hàm lượng oxy tham chiếu.",
            False,
        ),
        ("", False),
        ("XÁC NHẬN", True),
        ("Họ tên chuyên gia: ______________________   Đơn vị: ______________________", False),
        ("Ngày kiểm: ____/____/________            Chữ ký: ______________________", False),
        ("Kết luận chung (đạt / không đạt / đạt có điều kiện): ______________________", False),
        ("", False),
        (
            "Sau khi nhận lại file: python backend/scripts/review_import.py <file.xlsx> sẽ ghi kết quả vào "
            "review_status của từng rule/dòng và lưu nhật ký quyết định.",
            False,
        ),
    ]
    for text, bold in lines:
        ws.append([text])
        c = ws.cell(row=ws.max_row, column=1)
        c.font = Font(bold=bold, size=13 if bold and ws.max_row == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def sheet_thresholds(wb: Workbook, d: Data, sample: int, summary: list[list[Any]]) -> None:
    for sheet_name, code in THRESHOLD_TABLES:
        t = d.tables.get(code)
        if t is None:  # document not loaded (e.g. partial test database)
            continue
        rows = sample_rows(d.rows(code), sample)
        rules = d.rules_by_row(code)
        value_cols = [c for c in t.columns if c not in ("tt", "parameter", "unit", "equipment")]
        headers = (
            [
                "Mã dòng",
                "Nhãn trích dẫn",
                "TT",
                "Thông số (nguồn)",
                "Thông số → mã",
                "Đơn vị (nguồn)",
                "Đơn vị → mã",
                "Thiết bị / nhóm (đường dẫn)",
                "Thiết bị → mã",
                "Điều kiện áp dụng (đã tách)",
                "Ngoại lệ",
            ]
            + [f"Giá trị {c}" for c in value_cols]
            + ["Chú thích", "Vị trí nguồn", "Ô gốc"]
        )
        out = []
        for r in rows:
            eq = [
                c.get("text")
                for c in r.conditions
                if c.get("type") == "equipment_path" and c.get("text")
            ]
            eq_id = next(
                (x.equipment_type_id for x in rules.get(r.row_id, []) if x.equipment_type_id), ""
            )
            raw = " | ".join(f"{k}: {v}" for k, v in r.raw_cells.items() if v)
            out.append(
                [
                    r.row_id,
                    r.citation_label,
                    r.row_no,
                    r.cells.get("parameter", ""),
                    r.entities.get("parameter_id", ""),
                    r.cells.get("unit", ""),
                    r.entities.get("unit_id", ""),
                    " > ".join(eq) or r.cells.get("equipment", ""),
                    eq_id,
                    cond_text(r.conditions),
                    "; ".join(r.exceptions),
                ]
                + [value_text(r.values.get(c)) for c in value_cols]
                + [d.footnotes(r.footnote_refs), d.source_ref(r), raw]
            )
        ws = wb.create_sheet(sheet_name)
        n_data = len([r for r in d.rows(code) if not r.is_group_header])
        write_table_sheet(
            ws,
            [
                f"{d.docs[t.document_id].document_code} — {t.title or code}",
                f"Bảng có {n_data} dòng giá trị, {len(rules)} dòng sinh LimitRule; mẫu kiểm {len(out)} dòng "
                f"(cột header gốc: {' / '.join(' '.join(h for h in hr if h) for hr in t.header_rows)})",
            ],
            headers,
            out,
            [22, 40, 6, 30, 12, 12, 10, 34, 22, 40, 30] + [14] * len(value_cols) + [40, 14, 50],
        )
        summary.append([sheet_name, code, n_data, len(out), sum(len(v) for v in rules.values())])


def sheet_methods(wb: Workbook, d: Data, sample: int, summary: list[list[Any]]) -> None:
    ws = wb.create_sheet("3-Phương pháp")
    headers = [
        "Mã dòng",
        "Bảng",
        "Nhãn trích dẫn",
        "TT",
        "Thông số / mục (nguồn)",
        "Thông số → mã",
        "Phương pháp (nguồn)",
        "Phương pháp đã tách (mã · giai đoạn · thứ tự)",
        "Vị trí nguồn",
    ]
    out = []
    for code in METHOD_TABLES:
        t = d.tables.get(code)
        if not t:
            continue
        rows = sample_rows(d.rows(code), sample)
        for r in rows:
            label_col = next(
                (
                    r.cells[k]
                    for k in ("parameter", "item", "thong_so", "loai_mau")
                    if r.cells.get(k)
                ),
                "",
            )
            method_src = " | ".join(
                v
                for k, v in r.cells.items()
                if k not in ("tt", "stt", "parameter", "item", "thong_so", "loai_mau") and v
            )
            methods = "; ".join(f"{x['code']} · {x['stage']} · {x['order']}" for x in r.methods)
            out.append(
                [
                    r.row_id,
                    f"{d.docs[t.document_id].document_code} {t.title or code}",
                    r.citation_label,
                    r.row_no,
                    label_col,
                    r.entities.get("parameter_id", ""),
                    method_src,
                    methods,
                    d.source_ref(r),
                ]
            )
        summary.append(
            [
                "3-Phương pháp",
                code,
                len([x for x in d.rows(code) if not x.is_group_header]),
                len(rows),
                0,
            ]
        )
    write_table_sheet(
        ws,
        [
            "Bảng phương pháp (QCVN 40 PL1, QCVN 19 PL1, TT10 PL 4.1–4.3)",
            "Kiểm: mã phương pháp đúng, thứ tự (phương pháp trọng tài = thứ tự 1), "
            "giai đoạn (lấy mẫu / hiện trường / phòng thí nghiệm)",
        ],
        headers,
        out,
        [24, 30, 44, 6, 30, 12, 60, 50, 14],
    )


def sheet_tt10_tech(wb: Workbook, d: Data, summary: list[list[Any]]) -> None:
    ws = wb.create_sheet("3b-TT10 kỹ thuật")
    headers = ["Mã dòng", "Bảng", "Nhãn trích dẫn", "Nội dung dòng (cột: giá trị)", "Vị trí nguồn"]
    out = []
    for code in TT10_TECH_TABLES:
        t = d.tables.get(code)
        if not t:
            continue
        rows = d.rows(code)
        for r in rows:
            out.append(
                [
                    r.row_id,
                    t.title or code,
                    r.citation_label,
                    " | ".join(f"{k}: {v}" for k, v in r.cells.items() if v),
                    d.source_ref(r),
                ]
            )
        summary.append(["3b-TT10 kỹ thuật", code, len(rows), len(rows), 0])
    write_table_sheet(
        ws,
        [
            "TT 10/2021 — bảng kỹ thuật in-scope (thiết bị quan trắc tự động, giới hạn đánh giá "
            "số liệu, thiết bị đo trực tiếp)",
            "Kiểm: tách cột đúng, không mất/ghép ô",
        ],
        headers,
        out,
        [22, 40, 50, 90, 14],
    )


def sheet_sectors(wb: Workbook, d: Data, summary: list[list[Any]]) -> None:
    ws = wb.create_sheet("4-Ngành PL2")
    code = "QCVN40_2025.PL2"
    rows = [r for r in d.rows(code) if not r.is_group_header]
    headers = [
        "Mã dòng",
        "Nhãn trích dẫn",
        "TT",
        "Loại hình (nguồn)",
        "Loại hình → mã",
        "Mã ngành VSIC",
        "Thông số đặc trưng (nguồn)",
        "Thông số đặc trưng → mã",
        "Thông số có thể phát sinh (điều kiện: thông số)",
        "Vị trí nguồn",
    ]
    out = []
    for r in rows:
        ents = dict(r.entities)
        out.append(
            [
                r.row_id,
                r.citation_label,
                r.row_no,
                r.cells.get("sector", ""),
                ents.get("sector_id", ""),
                ", ".join(r.vsic_codes),
                r.cells.get("characteristic_parameters", ""),
                ", ".join(ents.get("characteristic_parameters", []) or [])
                if isinstance(ents.get("characteristic_parameters"), list)
                else json.dumps(ents.get("characteristic_parameters", ""), ensure_ascii=False),
                "; ".join(
                    f"{c.get('condition_text') or '(mọi trường hợp)'}: {c.get('parameters_raw')}"
                    f" → {', '.join(c.get('parameter_ids') or [])}"
                    for c in r.conditional_parameters
                ),
                d.source_ref(r),
            ]
        )
    summary.append(["4-Ngành PL2", code, len(rows), len(rows), 0])
    write_table_sheet(
        ws,
        [
            "QCVN 40:2025 — Phụ lục 2: loại hình sản xuất, mã ngành, thông số đặc trưng",
            "Kiểm: mã ngành đúng, mọi thông số đặc trưng được gán đúng mã, điều kiện phát sinh tách đúng",
        ],
        headers,
        out,
        [22, 40, 6, 40, 20, 18, 50, 40, 60, 12],
    )


def sheet_tef(wb: Workbook, d: Data, summary: list[list[Any]]) -> None:
    ws = wb.create_sheet("5-Hệ số TEF")
    headers = [
        "Mã dòng",
        "Bảng",
        "Nhãn trích dẫn",
        "TT",
        "Chất (nguồn)",
        "Tên chuẩn hóa",
        "TEF",
        "Vị trí nguồn",
    ]
    out = []
    for code in TEF_TABLES:
        t = d.tables.get(code)
        if t is None:
            continue
        rows = d.rows(code)
        for r in rows:
            name = r.cells.get("congener") or r.cells.get("pah") or ""
            coef = r.coefficient or {}
            out.append(
                [
                    r.row_id,
                    t.title or code,
                    r.citation_label,
                    r.row_no,
                    name,
                    coef.get("name", ""),
                    coef.get("tef", r.cells.get("tef", "")),
                    d.source_ref(r),
                ]
            )
        summary.append(["5-Hệ số TEF", code, len(rows), len(rows), 0])
    write_table_sheet(
        ws,
        ["QCVN 19:2024 — Phụ lục 2 (Dioxin/Furan) và Phụ lục 3 (PAH): hệ số độc tương đương"],
        headers,
        out,
        [22, 24, 40, 6, 36, 30, 10, 12],
    )


def sheet_relations(wb: Workbook, d: Data, summary: list[list[Any]]) -> None:
    ws = wb.create_sheet("6-Quan hệ")
    rels = (
        d.s.execute(
            select(m.DocumentRelation).order_by(
                m.DocumentRelation.source_document_id,
                m.DocumentRelation.type,
                m.DocumentRelation.id,
            )
        )
        .scalars()
        .all()
    )
    headers = [
        "Mã",
        "Văn bản nguồn",
        "Loại quan hệ",
        "Văn bản đích",
        "Phạm vi",
        "Phần bị thay thế/bãi bỏ",
        "Điều kiện",
        "Nguồn gốc",
        "Bằng chứng / nhãn trích dẫn",
    ]
    out = []
    for r in rels:
        tgt = d.docs.get(r.target_document_id) if r.target_document_id else None
        out.append(
            [
                f"rel:{r.id}",
                d.docs[r.source_document_id].document_code,
                r.type,
                (tgt.document_code if tgt else r.target_code or ""),
                r.scope,
                r.target_part or "",
                r.condition or "",
                "sổ đăng ký (PO)" if r.origin == "registry" else "trích từ văn bản",
                r.evidence or r.citation_label or "",
            ]
        )
    summary.append(["6-Quan hệ", "document_relations", len(out), len(out), 0])
    write_table_sheet(
        ws,
        [
            "Quan hệ pháp lý: PROMULGATES / REPLACES / REPEALS / REFERENCES (từ sổ đăng ký và "
            "trích tự động từ Lời nói đầu / Điều 2)",
            "Kiểm: đúng văn bản đích, đúng phạm vi (toàn bộ / một phần: bảng, chương), đúng điều kiện chuyển tiếp",
        ],
        headers,
        out,
        [10, 22, 16, 30, 10, 30, 40, 16, 50],
    )


def sheet_transitions(wb: Workbook, d: Data, summary: list[list[Any]]) -> None:
    ws = wb.create_sheet("7-Chuyển tiếp")
    trs = (
        d.s.execute(
            select(m.TransitionRule).order_by(
                m.TransitionRule.document_id, m.TransitionRule.rule_id
            )
        )
        .scalars()
        .all()
    )
    headers = [
        "Mã",
        "Văn bản",
        "Loại",
        "Nội dung (nguyên văn)",
        "Ngày trong văn bản (ISO)",
        "Mốc tương đối đã quy đổi",
        "Nhãn trích dẫn",
        "Vị trí nguồn",
    ]
    out = [
        [
            t.rule_id,
            d.docs[t.document_id].document_code,
            t.kind,
            t.text,
            ", ".join(t.dates),
            ", ".join(t.relative_dates),
            t.citation_label,
            d.source_ref(t),
        ]
        for t in trs
    ]
    summary.append(["7-Chuyển tiếp", "transition_rules", len(out), len(out), 0])
    write_table_sheet(
        ws,
        [
            "Điều khoản chuyển tiếp / lộ trình / hiệu lực (TT 06/2025, TT 45/2024, TT 10/2021)",
            "Kiểm: loại (transition / roadmap / effective / default_zone), ngày và mốc quy đổi "
            "đúng (mốc 'ngày Thông tư có hiệu lực' của TT45 = 01/07/2025 do PO khai báo)",
        ],
        headers,
        out,
        [34, 20, 14, 80, 22, 30, 44, 12],
    )


def sheet_facts(wb: Workbook, d: Data, summary: list[list[Any]]) -> None:
    ws = wb.create_sheet("8-Sự kiện đăng ký")
    facts = (
        d.s.execute(
            select(m.RegistryFact).order_by(m.RegistryFact.document_id, m.RegistryFact.fact_type)
        )
        .scalars()
        .all()
    )
    headers = [
        "Mã",
        "Văn bản",
        "Loại",
        "Giá trị",
        "Nguồn",
        "Bằng chứng trong văn bản",
        "Phạm vi",
        "Khai báo bởi / ngày",
        "Ghi chú hệ thống",
    ]
    out = [
        [
            f"fact:{f.id}",
            d.docs[f.document_id].document_code,
            f.fact_type,
            f.value,
            "do quản trị viên khai báo" if f.source == "admin_declared" else "trích từ văn bản",
            f.evidence or "",
            f.scope or "",
            f"{f.declared_by or ''} {f.declared_at.date() if f.declared_at else ''}".strip(),
            f.note or "",
        ]
        for f in facts
    ]
    summary.append(["8-Sự kiện đăng ký", "registry_facts", len(out), len(out), 0])
    write_table_sheet(
        ws,
        [
            "Số hiệu, ngày ký, ngày hiệu lực, người ký — kể cả các giá trị PO khai báo cho bản dự thảo (ADR-0004)",
            "Kiểm: đúng với bản ký / công báo; với bản dự thảo xác nhận giá trị khai báo",
        ],
        headers,
        out,
        [10, 22, 20, 30, 22, 40, 16, 30, 40],
    )


def sheet_mappings(wb: Workbook, d: Data, summary: list[list[Any]]) -> None:
    ws = wb.create_sheet("9-Ánh xạ từ điển")
    rules = d.s.execute(select(m.LimitRule)).scalars().all()
    params: dict[tuple[str, str, str, str], list[str]] = {}
    equips: dict[tuple[str, str], list[str]] = {}
    for r in rules:
        params.setdefault(
            (r.parameter_raw, r.parameter_id or "", r.unit_raw, r.unit_id or ""), []
        ).append(r.citation_label)
        if r.equipment_type_id:
            eq = next(
                (
                    c.get("text")
                    for c in reversed(r.conditions)
                    if c.get("type") == "equipment_path"
                ),
                "",
            )
            equips.setdefault((eq or "", r.equipment_type_id), []).append(r.citation_label)
    headers = [
        "Mã",
        "Loại ánh xạ",
        "Tên trong văn bản",
        "Mã chuẩn hóa",
        "Đơn vị (nguồn)",
        "Đơn vị → mã",
        "Số rule dùng",
        "Ví dụ trích dẫn",
    ]
    out = []
    for i, ((praw, pid, uraw, uid), labels) in enumerate(
        sorted(params.items(), key=lambda kv: (kv[0][1], kv[0][0]))
    ):
        out.append([f"map:param:{i}", "thông số", praw, pid, uraw, uid, len(labels), labels[0]])
    for i, ((eraw, eid), labels) in enumerate(sorted(equips.items(), key=lambda kv: kv[0][1])):
        out.append([f"map:equip:{i}", "thiết bị", eraw, eid, "", "", len(labels), labels[0]])
    summary.append(["9-Ánh xạ từ điển", "parameters/units/equipment", len(out), len(out), 0])
    write_table_sheet(
        ws,
        [
            "Ánh xạ tên trong văn bản → mã chuẩn hóa (từ điển v1). Một mã sai kéo theo mọi rule dùng mã đó.",
            "Kiểm: tên thông số/thiết bị được gán đúng mã (xem data/dictionaries/*.yaml); đơn vị đúng",
        ],
        headers,
        out,
        [16, 12, 50, 26, 16, 12, 10, 60],
    )


def sheet_monitoring(wb: Workbook, d: Data, settings: Settings, summary: list[list[Any]]) -> None:
    path = settings.data_dir / "derived" / "TT_10_2021" / "monitoring.json"
    ws = wb.create_sheet("10-Quan trắc TT10")
    headers = [
        "Mã",
        "Loại yêu cầu",
        "Lĩnh vực",
        "Tóm tắt (hệ thống)",
        "Nguyên văn",
        "Điều kiện / số liệu đã tách",
        "Thông số",
        "Nhãn trích dẫn",
    ]
    out = []
    if path.exists():
        for q in json.loads(path.read_text(encoding="utf-8")):
            out.append(
                [
                    q["requirement_id"],
                    q["requirement_type"],
                    q["domain"],
                    q.get("summary", ""),
                    q["text"],
                    cond_text(q.get("conditions", []), skip=()),
                    ", ".join(q.get("parameters", [])),
                    q["citation_label"],
                ]
            )
    summary.append(["10-Quan trắc TT10", "monitoring_requirements", len(out), len(out), 0])
    write_table_sheet(
        ws,
        [
            "Yêu cầu quan trắc trích từ TT 10/2021 (Đ3, Đ16–18, Đ22, Phụ lục 5, 6) — loại: "
            "method_precedence, sampling_condition, sampling_port, direct_measurement, cems, "
            "frequency_pointer",
            "Kiểm: đúng loại yêu cầu, đúng lĩnh vực (nước thải / khí thải / chung), số liệu tách "
            "đúng, có thiếu yêu cầu quan trọng nào không (ghi ở Ghi chú)",
        ],
        headers,
        out,
        [16, 20, 12, 50, 80, 50, 20, 50],
    )


def sheet_scenarios(wb: Workbook, d: Data, settings: Settings, summary: list[list[Any]]) -> None:
    """Applicability scenarios (data/rules/applicability_scenarios_v1.yaml) run live for the expert."""
    import yaml

    from app.applicability.engine import Assessor
    from app.applicability.models import AssessmentRequest

    ws = wb.create_sheet("12-Kịch bản áp dụng")
    headers = [
        "Mã",
        "Lĩnh vực",
        "Kịch bản",
        "Ngữ cảnh (JSON)",
        "Quy chuẩn ứng viên / trạng thái",
        "Cột",
        "Chuyển tiếp",
        "Thông số kiểm soát",
        "Số ngưỡng",
        "Ngưỡng tiêu biểu",
        "Cảnh báo",
        "Còn thiếu",
        "Rule đã áp dụng",
    ]
    out = []
    path = settings.data_dir / "rules" / "applicability_scenarios_v1.yaml"
    if path.exists():
        doc = yaml.safe_load(path.read_text(encoding="utf-8"))
        a = Assessor(d.s, settings)
        for sc in doc["scenarios"]:
            r = a.assess(AssessmentRequest(domain=sc["domain"], context=sc["context"]))
            sample = "; ".join(
                f"{x.parameter_raw[:30]} {x.zone}: {x.op} {x.value if x.value is not None else f'{x.min}–{x.max}'}"
                + (f" (O₂ {x.ref_o2:g}%)" if x.ref_o2 is not None else "")
                for x in r.limit_rules[:6]
            )
            out.append(
                [
                    sc["id"],
                    sc["domain"],
                    sc["title"],
                    json.dumps(sc["context"], ensure_ascii=False),
                    "; ".join(f"{c.document_code}: {c.status}" for c in r.candidates)
                    or "(chưa xác định)",
                    (r.effective_zone or "") + (" (mặc định)" if r.zone_defaulted else ""),
                    f"{r.transition.applies}: {r.transition.explanation}" if r.transition else "",
                    ", ".join(r.parameter_set),
                    len(r.limit_rules),
                    sample,
                    "\n".join(r.warnings),
                    "\n".join(f"{mi.label}: {mi.reason}" for mi in r.missing_information),
                    ", ".join(x.rule_id for x in r.applied_rules if x.outcome == "applied"),
                ]
            )
    summary.append(["12-Kịch bản áp dụng", "applicability_scenarios_v1", len(out), len(out), 0])
    write_table_sheet(
        ws,
        [
            "Kịch bản xác định quy chuẩn áp dụng (engine chạy trực tiếp khi sinh hồ sơ) — PLAN.md §7, §13 Phase 3",
            "Kiểm: quy chuẩn ứng viên, cột, chuyển tiếp, tập thông số và ngưỡng tiêu biểu có đúng với văn bản "
            "không; kết luận 'còn thiếu' có hợp lý không; bổ sung kịch bản còn thiếu ở Ghi chú",
        ],
        headers,
        out,
        [10, 10, 40, 50, 30, 10, 50, 40, 8, 60, 50, 40, 30],
    )


def sheet_issues(wb: Workbook) -> None:
    ws = wb.create_sheet("11-Phát hiện")
    headers = ["Văn bản / chủ đề", "Vấn đề", "Cần chuyên gia / PO quyết định"]
    write_table_sheet(
        ws,
        ["Điểm hệ thống chưa xử lý được hoặc cần quyết định"],
        headers,
        [list(x) for x in KNOWN_ISSUES],
        [24, 80, 60],
    )


def sheet_summary(wb: Workbook, summary: list[list[Any]]) -> None:
    ws = wb.create_sheet("1-Tổng hợp", 1)
    ws.append(
        [
            "Sheet",
            "Bảng / nguồn",
            "Tổng dòng",
            "Dòng trong mẫu kiểm",
            "LimitRule liên quan",
            "Đúng",
            "Sai",
            "Cần làm rõ",
            "Chưa kiểm",
        ]
    )
    style_header(ws, 1, 9)
    # one line per sheet (a sheet may hold several tables); the COUNTIF covers the whole sheet
    grouped: dict[str, list[Any]] = {}
    for sheet, code, total, sampled, n_rules in summary:
        g = grouped.setdefault(sheet, [sheet, [], 0, 0, 0])
        g[1].append(code)
        g[2] += total
        g[3] += sampled
        g[4] += n_rules
    for sheet, codes, total, sampled, n_rules in grouped.values():
        ws.append([sheet, ", ".join(codes), total, sampled, n_rules])
        r = ws.max_row
        meta = SHEET_META.get(sheet)
        if meta and meta[2] >= meta[1]:
            col, first, last = meta
            ref = f"'{sheet}'!{col}{first}:{col}{last}"
            ws.cell(row=r, column=6, value=f'=COUNTIF({ref},"Đúng")')
            ws.cell(row=r, column=7, value=f'=COUNTIF({ref},"Sai")')
            ws.cell(row=r, column=8, value=f'=COUNTIF({ref},"Cần làm rõ")')
            ws.cell(row=r, column=9, value=f"=D{r}-F{r}-G{r}-H{r}")
    ws.append([])
    ws.append(
        [
            "Tổng",
            "",
            f"=SUM(C2:C{ws.max_row - 1})",
            f"=SUM(D2:D{ws.max_row - 1})",
            f"=SUM(E2:E{ws.max_row - 1})",
            f"=SUM(F2:F{ws.max_row - 1})",
            f"=SUM(G2:G{ws.max_row - 1})",
            f"=SUM(H2:H{ws.max_row - 1})",
            f"=SUM(I2:I{ws.max_row - 1})",
        ]
    )
    for c in range(1, 10):
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
    for i, w in enumerate([22, 34, 12, 18, 18, 10, 10, 14, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def build(out: Path, sample: int, settings: Settings) -> Path:
    today = dt.date.today().isoformat()
    maker = get_sessionmaker(settings)
    wb = Workbook()
    summary: list[list[Any]] = []
    with maker() as session:
        d = Data(session)
        sheet_guide(wb, today, sample)
        sheet_thresholds(wb, d, sample, summary)
        sheet_methods(wb, d, sample, summary)
        sheet_tt10_tech(wb, d, summary)
        sheet_sectors(wb, d, summary)
        sheet_tef(wb, d, summary)
        sheet_relations(wb, d, summary)
        sheet_transitions(wb, d, summary)
        sheet_facts(wb, d, summary)
        sheet_mappings(wb, d, summary)
        sheet_monitoring(wb, d, settings, summary)
        sheet_scenarios(wb, d, settings, summary)
        sheet_issues(wb)
    sheet_summary(wb, summary)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=None)
    ap.add_argument("--sample", type=int, default=30)
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    settings = Settings()
    out = (
        Path(args.out)
        if args.out
        else settings.data_dir / "review" / f"review-pack-v1-{dt.date.today().isoformat()}.xlsx"
    )
    path = build(out, args.sample, settings)
    print("written:", path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
