"""Excel export/import of the gold set for a future expert review and for the blind second
check (plan §2.5). Same look as the expert review pack; the reviewer fills the yellow columns."""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

from app.evaluation.gold_schema import GoldItem
from app.review.excel import REVIEW_COLS, write_table_sheet

SECOND_COLS = ["Đáp án của bạn", "Số liệu", "Nhãn trích dẫn", "Phải từ chối? (Có/Không)", "Ghi chú"]
DECISION = {"Đúng": "reviewed", "Sai": "rejected", "Cần làm rõ": "draft"}


@dataclass
class Decision:
    item_id: str
    result: str | None
    correction: str | None
    note: str | None
    answer: str | None = None
    numbers: str | None = None
    labels: str | None = None
    must_abstain: str | None = None


def _num_text(it: GoldItem) -> str:
    return "; ".join(
        f"{n.display or n.value} {n.unit or ''} [{n.role}]".strip() for n in it.expected_numbers
    )


def _cond_text(it: GoldItem) -> str:
    return (
        json.dumps(it.expected_conditions.axes(), ensure_ascii=False)
        if it.expected_conditions.axes()
        else ""
    )


def _labels(it: GoldItem) -> str:
    return "\n".join(f"[{r.grade}] {r.citation_label}" for r in it.relevant)


def _source(it: GoldItem) -> str:
    if not it.source_ref:
        return ""
    r = it.source_ref
    loc = f"{r.document_id} dòng {r.line_start}–{r.line_end}"
    if r.printed_page:
        loc += f", tr. {r.printed_page}"
    if r.table_code:
        loc += f" · {r.table_code} dòng {r.row_no} cột {'/'.join(r.column_path or [])}"
    q = it.verification.source_quote or ""
    return f"{loc}\n{q}"


def export_review(items: list[GoldItem], out: Path, *, mode: str = "expert") -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "0-Hướng dẫn"
    guide = [
        "Bộ câu hỏi vàng v2 — hồ sơ kiểm",
        "",
        "Chỉ điền các cột màu vàng. Cột KẾT QUẢ chọn: Đúng / Sai / Cần làm rõ.",
        "Đúng = câu hỏi hợp lý và đáp án + trích dẫn đúng theo văn bản;",
        "Sai = ghi giá trị/đáp án đúng ở cột 'Sửa thành';",
        "Cần làm rõ = câu hỏi mơ hồ hoặc cần chuyên gia khác; ghi lý do ở Ghi chú.",
        "Mọi số liệu phải đối chiếu với văn bản gốc tại vị trí ghi ở cột 'Nguồn'.",
    ]
    if mode == "second-check":
        guide += [
            "",
            "Chế độ kiểm mù: đáp án của tác giả bị ẩn. Hãy tự viết đáp án ngắn",
            "(số + đơn vị + điều kiện), số liệu, nhãn trích dẫn (Mục/Bảng/TT/Cột)",
            "và Phải từ chối? theo cách bạn đọc văn bản.",
        ]
    for g in guide:
        ws0.append([g])
    ws0.column_dimensions["A"].width = 120

    ws = wb.create_sheet("2-Câu hỏi")
    if mode == "second-check":
        headers = ["Mã", "Tầng", "Split", "Câu hỏi", "Thời điểm áp dụng", "Nguồn (vị trí)"]
        rows = []
        for it in items:
            loc = _source(it).split("\n")[0]
            rows.append(
                [
                    it.id,
                    it.stratum or "",
                    it.split or "",
                    it.question,
                    str(it.as_of_date or ""),
                    loc,
                ]
            )
        write_table_sheet(
            ws,
            ["Kiểm mù — câu hỏi"],
            headers,
            rows,
            [14, 16, 8, 60, 14, 40],
            review_cols=SECOND_COLS,
        )
    else:
        headers = [
            "Mã",
            "Tầng",
            "Split",
            "Câu hỏi",
            "Thời điểm áp dụng",
            "Đáp án kỳ vọng",
            "Số liệu kỳ vọng",
            "Điều kiện kỳ vọng",
            "Phải từ chối",
            "Cảnh báo kỳ vọng",
            "Nhãn trích dẫn (mức)",
            "Nguồn + trích nguyên văn",
            "Ghi chú tác giả",
        ]
        rows = []
        for it in items:
            rows.append(
                [
                    it.id,
                    it.stratum or "",
                    it.split or "",
                    it.question,
                    str(it.as_of_date or ""),
                    it.expected_answer,
                    _num_text(it),
                    _cond_text(it),
                    "Có" if it.must_abstain else "Không",
                    ", ".join(it.expected_warnings),
                    _labels(it),
                    _source(it),
                    it.notes,
                ]
            )
        write_table_sheet(
            ws,
            ["Câu hỏi và đáp án kỳ vọng"],
            headers,
            rows,
            [14, 16, 8, 50, 12, 45, 26, 26, 10, 18, 45, 50, 30],
            review_cols=REVIEW_COLS,
        )

    traps = [it for it in items if it.trap_type or it.must_abstain]
    if traps and mode != "second-check":
        wt = wb.create_sheet("3-Câu bẫy")
        headers = [
            "Mã",
            "Loại bẫy",
            "Câu hỏi",
            "Hành vi đúng (mode)",
            "Không được",
            "Cảnh báo kỳ vọng",
            "Ghi chú judge",
        ]
        rows = [
            [
                it.id,
                it.trap_type or "abstain",
                it.question,
                it.expected_behaviour.mode,
                ", ".join(it.expected_behaviour.must_not),
                ", ".join(it.expected_warnings),
                it.judge_notes or "",
            ]
            for it in traps
        ]
        write_table_sheet(
            wt, ["Câu bẫy / phải từ chối"], headers, rows, [14, 18, 60, 14, 30, 20, 40]
        )
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def read_review(path: Path) -> list[Decision]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb["2-Câu hỏi"]
    header_row = None
    for r in range(1, 10):
        if ws.cell(row=r, column=1).value == "Mã":
            header_row = r
            break
    if header_row is None:
        raise ValueError("header row 'Mã' not found")
    headers = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}
    out: list[Decision] = []
    for r in range(header_row + 1, ws.max_row + 1):
        iid = ws.cell(row=r, column=1).value
        if not iid:
            continue

        def g(name: str, row: int = r):
            c = idx.get(name)
            v = ws.cell(row=row, column=c).value if c else None
            return str(v).strip() if v not in (None, "") else None

        out.append(
            Decision(
                item_id=str(iid),
                result=g("KẾT QUẢ"),
                correction=g("Sửa thành / giá trị đúng"),
                note=g("Ghi chú chuyên gia") or g("Ghi chú"),
                answer=g("Đáp án của bạn"),
                numbers=g("Số liệu"),
                labels=g("Nhãn trích dẫn"),
                must_abstain=g("Phải từ chối? (Có/Không)"),
            )
        )
    return out


def apply_review(items: list[GoldItem], decisions: list[Decision], reviewer: str) -> int:
    by_id = {it.id: it for it in items}
    n = 0
    for d in decisions:
        it = by_id.get(d.item_id)
        if it is None or not d.result:
            continue
        status = DECISION.get(d.result)
        if status is None:
            continue
        it.status = status
        it.reviewer = reviewer
        if d.correction:
            it.verification.expert_correction = d.correction
        if d.note:
            it.notes = (it.notes + " | " if it.notes else "") + f"[chuyên gia] {d.note}"
        n += 1
    return n
