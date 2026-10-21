"""Query-frame labels of the gold set and the metrics of query understanding (plan
docs/plans/2026-10-12-g3-hieu-cau-hoi-v2.md, sections 3 and 4).

Labels say what the *question* states, not what the answer needs:

- items anchored to one table cell (T1, T2, B1, a trap about a limit) take their slots from the
  rule of that cell (`slots_by="cell"`); a multi-hop pattern tells which of them the question does
  not state (the column of an unzoned facility) and which second intent it carries;
- hand-written prose items take their intents from the sheet reviewed by the PO and their slots
  from a second sheet (`slots_by="draft"` until reviewed; drafts are left out of slot metrics).

Nothing here calls an LLM."""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.evaluation import stats
from app.evaluation.gold_resolver import Resolver
from app.evaluation.gold_schema import FrameSlot, GoldFrame, GoldItem

# second intent and question-level facts of each multi-hop pattern (tag `pattern_<x>`)
PATTERNS: dict[str, dict[str, Any]] = {
    "dz": {"intents": ["applicability"], "drop": ["zone"]},  # column not stated: default column
    "zd": {"intents": []},  # column described by its definition, resolved to the code
    "tr": {"intents": ["applicability"], "date": "2032-01-01"},
    "rp": {"intents": ["effectivity"]},
    "fn_toc": {"intents": ["applicability"]},
    "fn_o2": {"intents": ["definition"]},
}
# dimensions a cell-anchored label is complete for; other dimensions wait for a reviewed label
CELL_DIMENSIONS = ("parameter", "zone", "flow", "equipment", "fuel_type", "document", "date")
FLOW_SPLIT = 2000.0  # m3/day, the two flow bands of the wastewater table

_QCVN = re.compile(r"qcvn\s*(\d+)(-mt)?\s*:\s*(\d{4})", re.I)
_CIRCULAR = re.compile(
    r"(?:thông\s+tư|tt)\s*(?:số\s*)?(\d{1,3})/(\d{4})|(\d{1,3})/(\d{4})/tt-", re.I
)


def document_key(code: str) -> str | None:
    """One key for the ways a document is written: 'QCVN 40:2025', 'QCVN 40:2025/BTNMT' →
    'QCVN40:2025'; 'Thông tư 06/2025', '06/2025/TT-BTNMT' → 'TT6/2025'."""
    keys = document_keys(code)
    return keys[0] if keys else None


def document_keys(text: str) -> list[str]:
    found: list[tuple[int, str]] = []
    for m in _QCVN.finditer(text):
        found.append((m.start(), f"QCVN{int(m.group(1))}{(m.group(2) or '').upper()}:{m.group(3)}"))
    for m in _CIRCULAR.finditer(text):
        no, year = (m.group(1), m.group(2)) if m.group(1) else (m.group(3), m.group(4))
        found.append((m.start(), f"TT{int(no)}/{year}"))
    out: list[str] = []
    for _, key in sorted(found):
        if key not in out:
            out.append(key)
    return out


def flow_band(value: float) -> str:
    return "F<=2000" if value <= FLOW_SPLIT else "F>2000"


def _pattern(item: GoldItem) -> dict[str, Any]:
    for tag in item.tags:
        if tag.startswith("pattern_"):
            return PATTERNS.get(tag.removeprefix("pattern_"), {"intents": []})
    return {"intents": []}


def cell_slots(item: GoldItem, resolver: Resolver) -> list[FrameSlot] | None:
    """Slots of an item anchored to one limit rule; None when it has no rule."""
    if not item.target_rule_codes:
        return None
    rule = resolver.rule(item.target_rule_codes[0])
    if rule is None:
        return None
    pattern = _pattern(item)
    cond = item.expected_conditions
    pairs: list[tuple[str, str | None]] = [
        ("parameter", rule.get("parameter_id")),
        ("zone", cond.zone),
        ("flow", cond.flow_band),
        ("equipment", rule.get("equipment_type_id")),
        ("fuel_type", cond.fuel),
        ("date", pattern.get("date")),
    ]
    pairs += [("document", key) for key in document_keys(item.question)]
    drop = set(pattern.get("drop", []))
    return [FrameSlot(dimension=d, value=v) for d, v in pairs if v and d not in drop]


def template_frame(item: GoldItem, resolver: Resolver) -> GoldFrame | None:
    """Frame of a template item (T1, T2): intents from the pattern, slots from the cell."""
    slots = cell_slots(item, resolver)
    if slots is None:
        return None
    return GoldFrame(intents=["threshold", *_pattern(item)["intents"]], slots=slots)


# ---------------------------------------------------------------- prediction and score


@dataclass
class PredictedFrame:
    domain: str
    intents: list[str]
    slots: set[tuple[str, str, bool]] = field(default_factory=set)
    needs_llm: bool = False


def from_understanding(qu: Any, document_codes: dict[str, str]) -> PredictedFrame:
    """The frame the rule-based understanding in production amounts to (one intent, one value per
    condition, no equipment)."""
    slots: set[tuple[str, str, bool]] = {("parameter", p, False) for p in qu.parameters}
    if qu.zone:
        slots.add(("zone", qu.zone, False))
    if qu.flow_m3_day is not None:
        slots.add(("flow", flow_band(qu.flow_m3_day), False))
    if qu.fuel_type:
        slots.add(("fuel_type", qu.fuel_type, False))
    if qu.receiving_body:
        slots.add(("receiving_body", qu.receiving_body, False))
    for d in qu.document_ids:
        key = document_key(document_codes.get(d, d))
        if key:
            slots.add(("document", key, False))
    return PredictedFrame(
        domain=qu.domain, intents=[qu.intent], slots=slots, needs_llm=qu.needs_llm
    )


def _f1(tp: int, fp: int, fn: int) -> float | None:
    if tp + fp + fn == 0:
        return None
    return 2 * tp / (2 * tp + fp + fn)


def score_item(item: GoldItem, pred: PredictedFrame) -> dict[str, Any]:
    """One row of the understanding report. `both` is the gate quantity of D55: right domain and
    right main intent."""
    frame = item.frame
    assert frame is not None
    gold_intents = list(frame.intents)
    domain_ok = pred.domain == item.domain
    main_ok = bool(pred.intents) and pred.intents[0] == gold_intents[0]
    gi, pi = set(gold_intents), set(pred.intents)
    row: dict[str, Any] = {
        "item_id": item.id,
        "split": item.split,
        "group": item.group,
        "domain_ok": float(domain_ok),
        "intent_main_ok": float(main_ok),
        "both": float(domain_ok and main_ok),
        "intent_f1": _f1(len(gi & pi), len(pi - gi), len(gi - pi)),
        "needs_llm": float(pred.needs_llm),
        "gold": {"domain": item.domain, "intents": gold_intents},
        "pred": {"domain": pred.domain, "intents": pred.intents},
        "slots": None,
    }
    if frame.slots_by == "draft":
        return row
    dims = None if frame.slots_by == "reviewed" else set(CELL_DIMENSIONS)
    gold = {(s.dimension, s.value, s.negated) for s in frame.slots}
    got = {s for s in pred.slots if dims is None or s[0] in dims}
    by_dim: dict[str, list[int]] = defaultdict(lambda: [0, 0, 0])
    for s in gold & got:
        by_dim[s[0]][0] += 1
    for s in got - gold:
        by_dim[s[0]][1] += 1
    for s in gold - got:
        by_dim[s[0]][2] += 1
    row["slots"] = {d: tuple(v) for d, v in by_dim.items()}
    row["slot_errors"] = {"extra": sorted(got - gold), "missed": sorted(gold - got)}
    return row


def summarize(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Rates with confidence intervals, micro-F1 over slots, by group and by dimension."""

    def block(rs: list[dict[str, Any]]) -> dict[str, Any]:
        out: dict[str, Any] = {"n": len(rs)}
        for key in ("domain_ok", "intent_main_ok", "both", "intent_f1", "needs_llm"):
            out[key] = stats.summarize_metric([r[key] for r in rs])
        tot: dict[str, list[int]] = defaultdict(lambda: [0, 0, 0])
        for r in rs:
            for d, (tp, fp, fn) in (r["slots"] or {}).items():
                for i, v in enumerate((tp, fp, fn)):
                    tot[d][i] += v
                    tot["*"][i] += v
        out["slot_items"] = sum(1 for r in rs if r["slots"] is not None)
        out["slot_f1"] = _f1(*tot["*"]) if "*" in tot else None
        out["slot_by_dimension"] = {
            d: {"tp": v[0], "fp": v[1], "fn": v[2], "f1": _f1(*v)}
            for d, v in sorted(tot.items())
            if d != "*"
        }
        return out

    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for r in rows:
        groups[f"split:{r['split']}"].append(r)
        groups[f"group:{r['group']}"].append(r)
    return {"all": block(rows)} | {k: block(v) for k, v in sorted(groups.items())}


def render_markdown(title: str, summary: dict[str, Any], rows: list[dict[str, Any]]) -> str:
    def rate(m: dict[str, Any]) -> str:
        if not m.get("n") or m.get("mean") is None:
            return "–"
        ci = m.get("wilson95") or m.get("ci95")
        return f"{m['mean']:.3f} ({ci[0]:.2f}–{ci[1]:.2f})" if ci else f"{m['mean']:.3f}"

    def f1(v: float | None) -> str:
        return "–" if v is None else f"{v:.3f}"

    lines = [
        f"# {title}",
        "",
        "| Nhóm | Số câu | Đúng lĩnh vực | Đúng ý định chính | Đúng cả hai | F1 tập ý định"
        " | F1 slot (số câu có nhãn slot) | Quy tắc báo không chắc |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for name, b in summary.items():
        lines.append(
            f"| {name} | {b['n']} | {rate(b['domain_ok'])} | {rate(b['intent_main_ok'])}"
            f" | {rate(b['both'])} | {f1(b['intent_f1'].get('mean'))}"
            f" | {f1(b['slot_f1'])} ({b['slot_items']}) | {f1(b['needs_llm'].get('mean'))} |"
        )
    lines += ["", "## Slot theo chiều (toàn bộ)", "", "| Chiều | Đúng | Thừa | Thiếu | F1 |"]
    lines.append("|---|---|---|---|---|")
    for d, v in summary["all"]["slot_by_dimension"].items():
        lines.append(f"| {d} | {v['tp']} | {v['fp']} | {v['fn']} | {f1(v['f1'])} |")
    wrong = [r for r in rows if not r["both"]]
    lines += ["", f"## Câu sai lĩnh vực hoặc ý định chính ({len(wrong)})", ""]
    lines += ["| Mã | Nhóm | Nhãn | Hệ thống |", "|---|---|---|---|"]
    for r in wrong:
        g, p = r["gold"], r["pred"]
        lines.append(
            f"| {r['item_id']} | {r['group']} | {g['domain']} · {', '.join(g['intents'])}"
            f" | {p['domain']} · {', '.join(p['intents'])} |"
        )
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------- review sheets

INTENT_SHEET = "Nhãn ý định"
INTENT_HEADERS = ("Ý định chính", "Ý định phụ 1", "Ý định phụ 2")
# the document slot in a workbook that already has a column «Văn bản» for the legal basis
DOCUMENT_SLOT_HEADER = "Văn bản nêu trong câu hỏi"
SLOT_SHEET = "Nhãn slot"


def read_intent_sheet(path: Path, sheet: str = INTENT_SHEET) -> dict[str, list[str]]:
    """Intents by item id: main intent first, then up to two secondary ones. Columns are found
    by their header, so the review sheet and an authoring workbook are read the same way."""
    from openpyxl import load_workbook

    out: dict[str, list[str]] = {}
    ws = load_workbook(path, data_only=True)[sheet]
    header = [str(c.value or "").strip() for c in ws[1]]
    cols = [header.index(h) for h in INTENT_HEADERS if h in header]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        intents = [str(row[j]).strip() for j in cols if row[j] and str(row[j]).strip()]
        if intents:
            out[str(row[0]).strip()] = list(dict.fromkeys(intents))
    return out


# (key, header, width); a cell holds codes separated by ";", a leading "!" marks a negated value
SLOT_COLUMNS: list[tuple[str, str, int]] = [
    ("parameter", "Thông số", 16),
    ("zone", "Cột", 8),
    ("flow", "Dải lưu lượng", 12),
    ("equipment", "Thiết bị", 30),
    ("fuel_type", "Nhiên liệu", 14),
    ("receiving_body", "Nguồn tiếp nhận", 14),
    ("sector", "Ngành", 16),
    ("document", "Văn bản", 22),
    ("date", "Ngày / mốc", 14),
    ("measured_value", "Giá trị đo", 14),
]
SLOT_EXTRA = [("missing", "Dữ kiện thiếu", 22), ("sub_questions", "Câu con (lĩnh vực)", 22)]


def write_slot_sheet(
    path: Path, rows: list[tuple[GoldItem, PredictedFrame]], catalogue: dict[str, list[str]]
) -> None:
    """Sheet for the PO: one row per hand-written item without a table cell, slots pre-filled
    from the rule-based matcher (a draft to correct, not a label)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = SLOT_SHEET
    bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="top")
    head, fill = PatternFill("solid", fgColor="D9E1F2"), PatternFill("solid", fgColor="FFF8D6")
    cols = [("id", "Mã", 11), ("group", "Nhóm", 7), ("question", "Câu hỏi", 70)]
    cols += SLOT_COLUMNS + SLOT_EXTRA + [("note", "Ghi chú", 30)]
    for j, (_, title, width) in enumerate(cols, start=1):
        c = ws.cell(1, j, title)
        c.font, c.fill, c.alignment = bold, head, wrap
        ws.column_dimensions[get_column_letter(j)].width = width
    for r, (item, pred) in enumerate(rows, start=2):
        values: dict[str, str] = {
            "id": item.id,
            "group": item.group or "",
            "question": item.question,
        }
        for dim, _, _ in SLOT_COLUMNS:
            got = sorted(("!" if neg else "") + v for d, v, neg in pred.slots if d == dim)
            values[dim] = "; ".join(got)
        if item.frame and len(item.frame.intents) > 1 and "comparison" in item.frame.intents:
            values["sub_questions"] = (
                "wastewater; emission" if item.domain == "cross_domain" else ""
            )
        values["missing"] = "; ".join(item.frame.missing) if item.frame else ""
        for j, (key, _, _) in enumerate(cols, start=1):
            c = ws.cell(r, j, values.get(key, ""))
            c.alignment = wrap
            if j > 3:
                c.fill = fill
    ws.freeze_panes = "D2"
    cat = wb.create_sheet("Danh mục")
    for j, (name, codes) in enumerate(catalogue.items(), start=1):
        c = cat.cell(1, j, name)
        c.font, c.fill = bold, head
        cat.column_dimensions[get_column_letter(j)].width = 44
        for i, code in enumerate(codes, start=2):
            cat.cell(i, j, code)
    wb.save(path)


def read_slot_sheet(path: Path, sheet: str = SLOT_SHEET) -> dict[str, dict[str, Any]]:
    """Reviewed slots by item id: {"slots": [FrameSlot], "missing": [...], "sub_questions": [...]}."""
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)[sheet]
    titles = {t: k for k, t, _ in SLOT_COLUMNS + SLOT_EXTRA}
    names = [str(c.value or "").strip() for c in ws[1]]
    if DOCUMENT_SLOT_HEADER in names:  # an authoring workbook: «Văn bản» there is the basis
        titles = {k: v for k, v in titles.items() if v != "document"}
        titles[DOCUMENT_SLOT_HEADER] = "document"
    header = [titles.get(n) for n in names]
    out: dict[str, dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        rec: dict[str, Any] = {"slots": [], "missing": [], "sub_questions": []}
        for key, cell in zip(header, row, strict=False):
            if not key or cell is None:
                continue
            parts = [p.strip() for p in str(cell).split(";") if p.strip()]
            if key in ("missing", "sub_questions"):
                rec[key] = parts
                continue
            for p in parts:
                neg = p.startswith("!")
                rec["slots"].append(FrameSlot(dimension=key, value=p.lstrip("! "), negated=neg))
        out[str(row[0]).strip()] = rec
    return out


def production_predictor(settings: Any, session: Any) -> Any:
    """question, as_of_date → PredictedFrame of the understanding that runs in production (rules
    only, no LLM call)."""
    from sqlalchemy import select

    from app.db import models as m
    from app.retrieval.query import understand
    from app.retrieval.retriever import RetrievalService

    svc = RetrievalService(settings)
    svc.ensure(session)
    codes = {d.document_id: d.document_code for d in session.execute(select(m.Document)).scalars()}

    def predict(question: str, as_of_date: Any = None) -> PredictedFrame:
        qu = understand(question, svc.aliases, as_of_date, svc.doc_domains)
        return from_understanding(qu, codes)

    return predict


def frame_predictor(settings: Any, session: Any) -> Any:
    """question → PredictedFrame of the rule path of query understanding v2 (no LLM call)."""
    from sqlalchemy import select

    from app.db import models as m
    from app.ontology.loader import load_dictionaries
    from app.retrieval.frame import FrameBuilder
    from app.retrieval.retriever import RetrievalService

    svc = RetrievalService(settings)
    svc.ensure(session)
    codes = {d.document_id: d.document_code for d in session.execute(select(m.Document)).scalars()}
    builder = FrameBuilder(
        load_dictionaries(settings=settings), svc.aliases, svc.corpus, svc.doc_domains
    )

    def predict(question: str, as_of_date: Any = None) -> PredictedFrame:
        return to_predicted(builder.build(question), codes)

    return predict


def to_predicted(frame: Any, codes: dict[str, str]) -> PredictedFrame:
    """A `QueryFrame` in the terms of the labels: flow as its band, a document as its key; a slot
    a hop filled is not something the question states."""
    slots: set[tuple[str, str, bool]] = set()
    for s in frame.slots:
        if s.source == "hop":
            continue
        value = s.value
        if s.dimension == "flow":
            value = flow_band(float(s.value))
        elif s.dimension == "document":
            value = document_key(codes.get(s.value, s.value)) or s.value
        slots.add((s.dimension, value, s.negated))
    return PredictedFrame(
        domain=frame.domain, intents=frame.intents, slots=slots, needs_llm=bool(frame.unsure)
    )


def frame_resolver(settings: Any, session: Any) -> tuple[Any, dict[str, str]]:
    """(FrameResolver of query understanding v2, document id → code): rules, candidate facts and
    hops; no LLM call."""
    from sqlalchemy import select

    from app.db import models as m
    from app.ontology.loader import load_dictionaries
    from app.retrieval.frame import FrameBuilder
    from app.retrieval.hops import FrameResolver
    from app.retrieval.retriever import RetrievalService

    svc = RetrievalService(settings)
    svc.ensure(session)
    d = load_dictionaries(settings=settings)
    codes = {x.document_id: x.document_code for x in session.execute(select(m.Document)).scalars()}
    builder = FrameBuilder(d, svc.aliases, svc.corpus, svc.doc_domains)
    return FrameResolver(session, d, builder, svc.corpus), codes
