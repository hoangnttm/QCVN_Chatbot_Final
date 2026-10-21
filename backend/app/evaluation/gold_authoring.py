"""Authoring kit of gold v3, round 1 (docs/plans/2026-10-06-xay-lai-bo-cau-hoi.md §0).

`build_workbook` writes the Excel file the author fills in: one pre-numbered row per quota slot of
`strata_v3.yaml`. Rows of group B1 already carry a table cell read from the raw source (the author
only writes the question); for the other groups the author gives the question, the short answer,
the document, the legal address, a verbatim quote and the expected behaviour.

`import_workbook` turns the filled rows into `GoldItem`s. Everything else (source lines, chunk
mapping, numbers, adjacent cells, facts, warnings) is derived here; a row is handed back to the
author when the quote is not in the source, the address contradicts the place the quote was found,
or the item does not pass the validator. Nothing in this module calls an LLM."""

from __future__ import annotations

import datetime as dt
import random
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.core.config import Settings
from app.evaluation.gold_candidates import candidate_chunks
from app.evaluation.gold_generate import DOMAIN_OF_TABLE, Generator
from app.evaluation.gold_resolver import DerivedResolver, Resolver
from app.evaluation.gold_schema import (
    ExpectedBehaviour,
    ExpectedNumber,
    GoldItem,
    Provenance,
    RelevantChunk,
    SourceRef,
    Strata,
    Stratum,
    Verification,
)
from app.evaluation.gold_source import (
    Derived,
    RawSource,
    nearest_line,
    normalize_raw,
    quote_matches,
)
from app.evaluation.gold_v3 import FactIndex
from app.evaluation.gold_validate import GoldValidator, dates_in
from app.evaluation.rawtables import RawCell
from app.qa.validator import numbers

AS_OF = dt.date(2026, 1, 15)
SHEET = "Soạn"
B1_TABLES = ("QCVN40_2025.B1", "QCVN40_2025.B2", "QCVN19_2024.B1", "QCVN19_2024.B2")

# (key, header, width) — the import reads columns by header text
COLUMNS: list[tuple[str, str, int]] = [
    ("id", "Mã", 11),
    ("group", "Nhóm", 7),
    ("topic", "Chủ đề", 34),
    ("split", "Tập", 6),
    ("hint", "Cho sẵn / gợi ý", 48),
    ("question", "Câu hỏi", 60),
    ("answer", "Đáp án ngắn", 36),
    ("document", "Văn bản", 22),
    ("address", "Địa chỉ pháp lý", 34),
    ("quote", "Trích nguyên văn", 50),
    ("behaviour", "Hành vi mong đợi", 16),
    ("domain", "Lĩnh vực", 14),
    ("note", "Ghi chú (dữ kiện thiếu / lý do từ chối)", 34),
    ("document2", "Văn bản 2", 22),
    ("address2", "Địa chỉ 2", 30),
    ("quote2", "Trích 2", 44),
    ("stratum", "stratum", 14),
    ("cell", "cell", 14),
    ("trap", "trap", 14),
]
AUTHOR_KEYS = ("question", "answer", "document", "address", "quote", "behaviour", "domain", "note")
HIDDEN = ("stratum", "cell", "trap")

BEHAVIOURS = {
    "trả lời": "answer",
    "từ chối": "abstain",
    "chỉ dẫn": "pointer",
    "hỏi lại": "clarify",
}
DOMAINS = {
    "nước thải": "wastewater",
    "khí thải": "emission",
    "cả hai": "cross_domain",
    "ngoài phạm vi": "out_of_scope",
}
DOMAIN_VI = {v: k for k, v in DOMAINS.items()}
INTENT = {
    "WW-MGMT": "other",
    "WW-SECTOR": "scope",
    "DEF-SCOPE": "definition",
    "SAMPLING": "sampling",
    "METHOD": "method",
    "EFFECTIVITY": "effectivity",
    "TRANSITION": "applicability",
    "XDOC": "other",
    "CLARIFY": "threshold",
    "NEG-OOS": "other",
    "FREQ-M01": "frequency",
    "NEG-MISSING": "other",
    "NEG-TRAP": "threshold",
}
DEFAULT_BEHAVIOUR = {
    "CLARIFY": "hỏi lại",
    "NEG-OOS": "từ chối",
    "FREQ-M01": "chỉ dẫn",
    "NEG-MISSING": "từ chối",
}
TRAPS = {
    "false_premise": ("tiền đề sai", "numbers_outside_evidence"),
    "pass_fail_request": ("hỏi đạt hay không đạt", "pass_fail_verdict"),
    "injection": ("câu hỏi chèn chỉ thị cho hệ thống", "follow_injected_instruction"),
    "unit_conversion": ("yêu cầu quy đổi đơn vị", "unit_conversion"),
    "replaced_standard": ("hỏi theo quy chuẩn đã bị thay thế", "use_replaced_values"),
}
TRAP_PLAN = {
    "dev": ["false_premise", "pass_fail_request"],
    "test": ["false_premise", "injection", "unit_conversion", "replaced_standard"],
}
STRATUM_TRAP = {"NEG-OOS": "out_of_scope", "NEG-MISSING": "missing_information"}

_ADDR_RE = re.compile(
    r"(điều|khoản|điểm|mục|bảng|phụ lục|chương|cột)\s+([0-9a-zđ]+(?:\.[0-9a-z]+)*)", re.I
)
_DATE_STRIP = re.compile(
    r"\b\d{1,2}[/.-]\d{1,2}[/.-]\d{4}\b|\b\d{4}-\d{2}-\d{2}\b"
    r"|ngày\s+\d{1,2}\s+tháng\s+\d{1,2}\s+năm\s+\d{4}",
    re.I,
)

# document codes and unit numbers ("Bảng 1", "QCVN 40:2025", "06/2025/TT-BTNMT") are not values
_REF_STRIP = re.compile(
    r"qcvn\s*\d+(?:-mt)?\s*:\s*\d{4}(?:/btnmt)?|\d+/\d{4}(?:/[a-zđ-]+)?"
    r"|(?:điều|khoản|điểm|mục|bảng|phụ lục|chương|cột|tt)\s+[0-9]+(?:\.[0-9a-z]+)*",
    re.I,
)


def _values(text: str) -> set[str]:
    return numbers(_REF_STRIP.sub(" ", _DATE_STRIP.sub(" ", text or "")))


GUIDE = [
    "HƯỚNG DẪN SOẠN CÂU HỎI — ĐỢT 1 (80 câu)",
    "",
    "1. Làm ở sheet «Soạn». Mỗi dòng là một câu đã có sẵn Mã, Nhóm, Chủ đề. Chỉ điền các cột nền vàng.",
    "2. Viết câu hỏi theo giọng người dùng thật (cán bộ môi trường của nhà máy, tư vấn, cán bộ sở)."
    " Không chép nguyên văn điều khoản; không chép lại câu ở sheet «Ví dụ».",
    "3. Đáp án và căn cứ đọc từ VĂN BẢN GỐC. Không hỏi chatbot, không xem kết quả của hệ thống.",
    "",
    "NHÓM B1 — tra ngưỡng: ô bảng, đáp án, địa chỉ đã điền sẵn ở cột «Cho sẵn». Bạn chỉ viết câu hỏi."
    " Câu hỏi phải nêu đủ điều kiện để chỉ ra đúng một ô: thông số, cột (hoặc mô tả vùng), và"
    " lưu lượng cụ thể nằm trong dải (QCVN 40 Bảng 1) hoặc thiết bị, nhiên liệu, công suất (QCVN 19)."
    " Không viết con số đáp án vào câu hỏi.",
    "NHÓM B2, B3 — văn xuôi, hiệu lực: điền Câu hỏi, Đáp án ngắn, Văn bản, Địa chỉ pháp lý, Trích nguyên văn.",
    "   • Địa chỉ pháp lý: ghi tới mức nhỏ nhất, ví dụ «Điều 2, Khoản 4», «Mục 4.3.2», «Phụ lục 2»;"
    " thêm số trang nếu có («tr. 12»).",
    "   • Trích nguyên văn: MỘT câu hoặc một vế chứa đáp án, chép đúng từng chữ từ văn bản"
    " (nên chép từ cột Nội dung của sheet «Gợi ý»). Không cần chép cả điều khoản.",
    "   • Ngày hiệu lực của Thông tư 45/2024 và Thông tư 10/2021 không in trong tệp: ghi Địa chỉ ="
    " «Sổ đăng ký», để trống ô Trích, Đáp án ghi ngày dạng dd/mm/yyyy.",
    "NHÓM B4 — liên văn bản: câu hỏi cần căn cứ ở HAI văn bản. Điền thêm Văn bản 2, Địa chỉ 2, Trích 2.",
    "NHÓM B5 — thiếu dữ kiện: câu hỏi cố ý thiếu một điều kiện (cột, lưu lượng, thiết bị…)."
    " Cột Ghi chú: liệt kê dữ kiện thiếu, cách nhau bằng dấu chấm phẩy."
    " Đáp án ngắn: nếu chỉ còn tối đa 3 phương án thì ghi cả các phương án; nhiều hơn thì ghi «hỏi lại»."
    " Có thể để trống Văn bản, Địa chỉ, Trích.",
    "NHÓM B6 — từ chối và bẫy:",
    "   • Ngoài phạm vi / không có trong corpus: Hành vi = «từ chối»; Ghi chú nêu lý do; không cần căn cứ.",
    "   • Tần suất quan trắc: Hành vi = «chỉ dẫn»; căn cứ là điều khoản của Thông tư 10/2021 chỉ sang"
    " pháp luật hiện hành (Điều 16 khoản 3, Điều 17 khoản 5).",
    "   • Bẫy: loại bẫy ghi ở cột Chủ đề. Nếu câu vẫn trả lời được (ví dụ hỏi «đạt hay không» thì hệ thống"
    " nêu ngưỡng nhưng không kết luận) thì Hành vi = «trả lời» và điền căn cứ như B2;"
    " nếu phải từ chối thì Hành vi = «từ chối».",
    "",
    "CỘT «Lĩnh vực»: nước thải / khí thải / cả hai / ngoài phạm vi. Để trống thì công cụ lấy theo văn bản;"
    " riêng Thông tư 10/2021 dùng cho cả hai lĩnh vực nên phải chọn.",
    "CỘT «Tập»: dev = dùng để chỉnh hệ thống, test = để kiểm tra cuối. Không đổi giá trị cột này.",
    "",
    "Sau mỗi lô, gửi tệp lại cho nhóm kỹ thuật. Công cụ sẽ trả về danh sách câu cần sửa kèm lý do"
    " (không tìm thấy đoạn trích, địa chỉ lệch với chỗ tìm thấy đoạn trích, câu hỏi trùng…).",
    "",
    "Sheet «Mẫu câu»: viết 15 mẫu câu mới cho phần sinh tự động (T1, T2). Giữ nguyên các chỗ trống"
    " trong ngoặc nhọn; chỉ đổi cách diễn đạt.",
]

EXAMPLES = [
    (
        "B4",
        "Địa phương chưa xác định phân vùng thì nước thải công nghiệp và khí thải công nghiệp"
        " áp dụng cột nào của hai quy chuẩn?",
        "Cột B của Bảng 1 và Bảng 2 trong cả QCVN 40:2025/BTNMT và QCVN 19:2024/BTNMT",
        "06/2025/TT-BTNMT",
        "Điều 2, Khoản 4",
        "trường hợp chưa xác định được phân vùng xả thải nước thải thì áp dụng Cột B quy định tại"
        " Bảng 1 và Bảng 2 của QCVN 40:2025/BTNMT",
        "trả lời",
        "Văn bản 2: 45/2024/TT-BTNMT · Địa chỉ 2: Điều 2, Khoản 4 · Trích 2: Trường hợp chưa xác định"
        " được phân vùng môi trường thì áp dụng Cột B quy định tại Bảng 1 và Bảng 2 của QCVN 19:2024/BTNMT",
    ),
    (
        "B5",
        "Nhà máy xả 1500 m3 nước thải mỗi ngày thì BOD5 tối đa là bao nhiêu?",
        "Cột A: 40 mg/L; Cột B: 60 mg/L; Cột C: 80 mg/L",
        "QCVN 40:2025/BTNMT",
        "Mục 2.1, Bảng 1",
        "",
        "hỏi lại",
        "Ghi chú: cột (phân vùng xả thải). Còn 3 phương án nên liệt kê cả ba.",
    ),
    (
        "B5",
        "Giới hạn COD trong nước thải công nghiệp là bao nhiêu?",
        "hỏi lại",
        "",
        "",
        "",
        "hỏi lại",
        "Ghi chú: lưu lượng xả thải; cột (phân vùng xả thải). Còn 6 phương án nên hỏi lại.",
    ),
]


# ---------------------------------------------------------------- B1 cells


def _is_cell_of(code: str, cell: RawCell) -> bool:
    if code == "WW-B1-FLOWZONE":
        return cell.table_code == "QCVN40_2025.B1"
    if code == "WW-B2-PLAIN":
        return cell.table_code == "QCVN40_2025.B2" and "." not in cell.row_key
    if code == "WW-B2-COND":
        return cell.table_code == "QCVN40_2025.B2" and "." in cell.row_key
    if not cell.table_code.startswith("QCVN19"):
        return False
    fallback = bool(cell.row_label_path) and "xả thải khác" in cell.row_label_path[-1].lower()
    if code == "EM-EXCL-FALLBACK":
        return fallback
    if fallback:
        return False
    if code == "EM-O2-FOOTNOTE":
        return cell.o2_ref is not None
    if code == "EM-B1-EQUIP":
        return bool(cell.equipment_path) and cell.equipment_path.count(".") >= 1
    return False


def _cluster(cell: RawCell) -> str:
    return f"{cell.table_code}.r{cell.row_key}"


def pick_cells(
    gen: Generator,
    stratum: Stratum,
    n_dev: int,
    n_test: int,
    *,
    seed: int,
    used_rules: set[str],
    dev_clusters: set[str],
    taken_rows: dict[str, str],
) -> list[tuple[str, RawCell]]:
    """Seeded choice of cells for one B1 stratum: zones rotated, parameters spread, cells of
    earlier hand-written items left out, and no test cell from a row whose template questions are
    in dev. One cell per table row; a table with fewer rows than the quota gives a second cell of
    a row already used in the same split."""
    pool: list[RawCell] = []
    for t in B1_TABLES:
        if t in stratum.tables:
            pool += [c for c in gen.reader.cells(t) if _is_cell_of(stratum.code, c)]
    pool = [c for c in pool if c.cell_key not in used_rules and gen._rule_for(c) is not None]
    rng = random.Random(f"{seed}:{stratum.code}")
    rng.shuffle(pool)
    out: list[tuple[str, RawCell]] = []
    params: dict[str, int] = defaultdict(int)
    for split, n in (("dev", n_dev), ("test", n_test)):
        got = 0
        while got < n:
            want = "ABC"[len(out) % 3]
            free = [
                c
                for c in pool
                if c.cell_key not in {x.cell_key for _, x in out}
                and taken_rows.get(_cluster(c), split) == split
                and (split == "dev" or _cluster(c) not in dev_clusters)
            ]
            ok = [c for c in free if _cluster(c) not in taken_rows] or free
            if not ok:
                break
            ok.sort(key=lambda c: (params[c.parameter], 0 if c.zone == want else 1))
            cell = ok[0]
            taken_rows[_cluster(cell)] = split
            params[cell.parameter] += 1
            out.append((split, cell))
            got += 1
    return out


def _cell_hint(cell: RawCell) -> str:
    parts = [f"Thông số: {cell.parameter}"]
    if len(cell.row_label_path) > 1:
        what = "Thiết bị" if cell.table_code.startswith("QCVN19") else "Trường hợp"
        parts.append(f"{what}: " + " › ".join(cell.row_label_path[1:]))
    if cell.flow_band:
        band = "F ≤ 2000" if cell.flow_band == "F<=2000" else "F > 2000"
        parts.append(f"Dải lưu lượng: {band} m³/ngày (nêu một lưu lượng cụ thể trong dải)")
    parts.append(f"Cột: {cell.zone}")
    if cell.o2_ref is not None:
        parts.append(f"Ôxy tham chiếu: {cell.o2_ref:g} %")
    return "\n".join(parts)


# ---------------------------------------------------------------- workbook


def _docs(raw: RawSource) -> dict[str, str]:
    """document_code → document_id for the documents that have a source file."""
    return {d.document_code: d.document_id for d in raw.registry.documents if d.source}


def build_workbook(
    settings: Settings, strata: Strata, existing: list[GoldItem], out: Path, *, seed: int | None
) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    seed = strata.seed if seed is None else seed
    gen = Generator(settings, settings.data_dir / "gold" / "templates_v2.yaml", seed)
    resolver = DerivedResolver(settings)
    raw = RawSource(settings)
    codes = {v: k for k, v in _docs(raw).items()}
    used_rules = {c for it in existing if it.tier == "B" for c in it.target_rule_codes}
    dev_clusters = {it.cluster_id for it in existing if it.split == "dev" and it.cluster_id}
    migrated = defaultdict(int)
    for it in existing:
        if it.tier == "B" and it.split == "dev":
            migrated[it.stratum] += 1

    rows: list[dict[str, Any]] = []
    hints: list[tuple[str, str, str, str]] = []
    counter: dict[str, int] = defaultdict(int)
    taken_rows: dict[str, str] = {}
    for s in strata.strata:
        if s.tier != "B" or not s.group:
            continue
        n_dev, n_test = max(0, s.quota_dev - migrated[s.code]), s.quota_test
        base = {"group": s.group, "topic": s.title, "stratum": s.code}
        if s.group == "B1":
            for split, cell in pick_cells(
                gen,
                s,
                n_dev,
                n_test,
                seed=seed,
                used_rules=used_rules,
                dev_clusters=dev_clusters,
                taken_rows=taken_rows,
            ):
                rule = gen._rule_for(cell) or {}
                rows.append(
                    base
                    | {
                        "split": split,
                        "hint": _cell_hint(cell),
                        "answer": gen._expected_answer(cell),
                        "document": codes.get(cell.document_id, cell.document_id),
                        "address": rule.get("citation_label", ""),
                        "quote": cell.display,
                        "behaviour": "trả lời",
                        "domain": DOMAIN_VI[DOMAIN_OF_TABLE[cell.table_code]],
                        "cell": cell.cell_key,
                    }
                )
            continue
        traps = {sp: list(TRAP_PLAN[sp]) for sp in TRAP_PLAN} if s.code == "NEG-TRAP" else {}
        for split, n in (("dev", n_dev), ("test", n_test)):
            for _ in range(n):
                row = base | {"split": split, "behaviour": DEFAULT_BEHAVIOUR.get(s.code, "")}
                if s.domain:
                    row["domain"] = DOMAIN_VI.get(s.domain, "")
                if traps.get(split):
                    row["trap"] = traps[split].pop(0)
                    row["topic"] = f"Bẫy — {TRAPS[row['trap']][0]}"
                rows.append(row)
        if s.chunk_ids or s.rule_ids or (s.document_ids and s.chunk_types):
            for c in candidate_chunks(resolver, s, seed=seed, n=max(8, 3 * s.quota)):
                hints.append(
                    (
                        s.title,
                        codes.get(c.document_id, c.document_id),
                        c.label,
                        re.sub(r"\s+", " ", c.text)[:700],
                    )
                )
    order = {g: i for i, g in enumerate(("B1", "B2", "B3", "B4", "B5", "B6"))}
    rows.sort(key=lambda r: order[r["group"]])
    for r in rows:
        counter[r["group"]] += 1
        r["id"] = f"G-{r['group']}-{counter[r['group']]:03d}"

    wb = Workbook()
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    fill_in = PatternFill("solid", fgColor="FFF8D6")
    head = PatternFill("solid", fgColor="D9E1F2")

    ws = wb.active
    ws.title = "Hướng dẫn"
    ws.column_dimensions["A"].width = 130
    for i, line in enumerate(GUIDE, start=1):
        cell = ws.cell(i, 1, line)
        cell.alignment = wrap
        if i == 1 or line.startswith(("NHÓM", "CỘT")):
            cell.font = bold

    ws = wb.create_sheet(SHEET)
    keys = [k for k, _, _ in COLUMNS]
    for j, (k, title, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(1, j, title)
        c.font, c.fill, c.alignment = bold, head, wrap
        col = get_column_letter(j)
        ws.column_dimensions[col].width = width
        ws.column_dimensions[col].hidden = k in HIDDEN
    for i, r in enumerate(rows, start=2):
        for j, k in enumerate(keys, start=1):
            c = ws.cell(i, j, r.get(k, ""))
            c.alignment = wrap
            fill_me = k in AUTHOR_KEYS or (
                k in ("document2", "address2", "quote2") and r["group"] == "B4"
            )
            if r["group"] == "B1":
                fill_me = k == "question"
            if fill_me:
                c.fill = fill_in
    ws.freeze_panes = "F2"
    last = len(rows) + 1
    lists = wb.create_sheet("Danh mục")
    for j, values in enumerate((list(_docs(raw)), list(BEHAVIOURS), list(DOMAINS)), start=1):
        for i, v in enumerate(values, start=1):
            lists.cell(i, j, v)
    for key, col_list, n in (
        ("document", "A", len(_docs(raw))),
        ("document2", "A", len(_docs(raw))),
        ("behaviour", "B", len(BEHAVIOURS)),
        ("domain", "C", len(DOMAINS)),
    ):
        dv = DataValidation(
            type="list", formula1=f"='Danh mục'!${col_list}$1:${col_list}${n}", allow_blank=True
        )
        col = get_column_letter(keys.index(key) + 1)
        dv.add(f"{col}2:{col}{last}")
        ws.add_data_validation(dv)

    ws = wb.create_sheet("Gợi ý")
    for j, (title, width) in enumerate(
        (("Chủ đề", 34), ("Văn bản", 22), ("Địa chỉ pháp lý", 44), ("Nội dung", 120)), start=1
    ):
        c = ws.cell(1, j, title)
        c.font, c.fill = bold, head
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, h in enumerate(hints, start=2):
        for j, v in enumerate(h, start=1):
            ws.cell(i, j, v).alignment = wrap
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Ví dụ")
    titles = (
        "Nhóm",
        "Câu hỏi",
        "Đáp án ngắn",
        "Văn bản",
        "Địa chỉ pháp lý",
        "Trích",
        "Hành vi",
        "Thêm",
    )
    for j, (title, width) in enumerate(
        zip(titles, (7, 60, 40, 22, 40, 50, 12, 60), strict=True), 1
    ):
        c = ws.cell(1, j, title)
        c.font, c.fill = bold, head
        ws.column_dimensions[get_column_letter(j)].width = width
    mode_vi = {v: k for k, v in BEHAVIOURS.items()}
    examples = [
        (
            it.group,
            it.question,
            it.expected_answer,
            codes.get(it.source_ref.document_id, "") if it.source_ref else "",
            "; ".join(it.legal_address),
            it.verification.source_quote or "",
            mode_vi[it.expected_behaviour.mode],
            "câu dev có sẵn — chỉ để tham khảo cách viết",
        )
        for it in existing
        if it.tier == "B"
    ] + [(*e,) for e in EXAMPLES]
    for i, e in enumerate(sorted(examples, key=lambda e: e[0] or ""), start=2):
        for j, v in enumerate(e, start=1):
            ws.cell(i, j, v).alignment = wrap

    ws = wb.create_sheet("Mẫu câu")
    _template_sheet(ws, gen, bold, head, wrap, fill_in)

    wb.save(out)
    return len(rows)


def _template_sheet(ws: Any, gen: Generator, bold: Any, head: Any, wrap: Any, fill_in: Any) -> None:
    import yaml

    spec = yaml.safe_load(
        (gen.settings.data_dir / "gold" / "multihop_v2.yaml").read_text(encoding="utf-8")
    )
    kinds: list[tuple[str, str, str, str]] = []
    t1 = gen.t["templates"]
    for key, what, n in (
        ("flow_zone", "T1 — QCVN 40 Bảng 1 (lưu lượng × cột)", 2),
        ("zone", "T1 — QCVN 40 Bảng 2 (cột)", 2),
        ("equipment_zone", "T1 — QCVN 19 (thiết bị × cột)", 2),
    ):
        old = t1[key][0]["text"]
        slots = ", ".join(sorted(set(re.findall(r"\{\w+\}", " ".join(x["text"] for x in t1[key])))))
        kinds += [(f"T1:{key}", what, slots, old)] * n
    names = {
        "DZ": "T2 — chưa xác định phân vùng",
        "ZD": "T2 — vùng mô tả bằng định nghĩa",
        "TR": "T2 — lộ trình 01/01/2032",
        "RP": "T2 — QCVN 40:2011 đã bị thay thế",
        "FN_TOC": "T2 — chọn TOC thay COD",
        "FN_O2": "T2 — ôxy tham chiếu",
    }
    domain_vi = {"wastewater": "nước thải", "emission": "khí thải"}
    for key, what in names.items():
        for domain, tpls in spec["templates"][key].items():
            slots = ", ".join(
                sorted(set(re.findall(r"\{\w+\}", " ".join(x["text"] for x in tpls))))
            )
            kinds.append(
                (f"T2:{key}:{domain}", f"{what} ({domain_vi[domain]})", slots, tpls[0]["text"])
            )
    for j, (title, width) in enumerate(
        (
            ("Kiểu", 16),
            ("Dùng cho", 38),
            ("Chỗ trống dùng được", 40),
            ("Mẫu cũ (tham khảo)", 70),
            ("Mẫu câu mới", 80),
        ),
        start=1,
    ):
        c = ws.cell(1, j, title)
        c.font, c.fill = bold, head
        ws.column_dimensions[chr(64 + j)].width = width
    for i, k in enumerate(kinds, start=2):
        for j, v in enumerate(k, start=1):
            ws.cell(i, j, v).alignment = wrap
        ws.cell(i, 5).fill = fill_in


def build_template_workbook(settings: Settings, out: Path) -> int:
    """Sheet «Mẫu câu» alone, in a file of its own (the main workbook may be open in Excel)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    gen = Generator(settings, settings.data_dir / "gold" / "templates_v2.yaml", 0)
    wb = Workbook()
    ws = wb.active
    ws.title = "Mẫu câu"
    _template_sheet(
        ws,
        gen,
        Font(bold=True),
        PatternFill("solid", fgColor="D9E1F2"),
        Alignment(wrap_text=True, vertical="top"),
        PatternFill("solid", fgColor="FFF8D6"),
    )
    wb.save(out)
    return ws.max_row - 1


def read_templates(path: Path) -> dict[str, list[str]]:
    """New template sentences of sheet «Mẫu câu» by kind (`T1:zone`, `T2:DZ:emission`, …). A
    sentence may only use the slots listed for its kind, and every kind needs at least one."""
    from openpyxl import load_workbook

    out: dict[str, list[str]] = defaultdict(list)
    problems: list[str] = []
    for kind, _, slots, _, new in load_workbook(path, data_only=True)["Mẫu câu"].iter_rows(
        min_row=2, max_col=5, values_only=True
    ):
        if not kind:
            continue
        text = str(new or "").strip()
        if not text:
            out.setdefault(kind, [])
            continue
        extra = set(re.findall(r"\{\w*\}", text)) - set(re.findall(r"\{\w+\}", slots or ""))
        if extra:
            problems.append(f"{kind}: chỗ trống không dùng được {sorted(extra)} trong {text!r}")
        out[kind].append(text)
    problems += [f"{kind}: chưa có mẫu câu mới" for kind, texts in out.items() if not texts]
    if any(k.count(":") == 1 for k in out if k.startswith("T2:")):
        problems.append("sheet «Mẫu câu» là bản cũ (T2 chưa tách nước thải / khí thải)")
    if problems:
        raise Reject("; ".join(problems))
    return dict(out)


def hand_cells(path: Path, settings: Settings) -> dict[str, dict[str, str]]:
    """B1 cells of the workbook: cell key → {split, row} (row = cluster id of the table row)."""
    from openpyxl import load_workbook

    gen = Generator(settings, settings.data_dir / "gold" / "templates_v2.yaml", 0)
    rows = {c.cell_key: _cluster(c) for t in B1_TABLES for c in gen.reader.cells(t)}
    ws = load_workbook(path, data_only=True)[SHEET]
    keys = [k for k, _, _ in COLUMNS]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(keys, r, strict=False))
        if d.get("cell"):
            out[d["cell"]] = {"split": d["split"], "row": rows[d["cell"]]}
    return out


# ---------------------------------------------------------------- import


@dataclass
class ImportResult:
    n_rows: int = 0
    items: list[GoldItem] = field(default_factory=list)
    rejects: dict[str, list[str]] = field(default_factory=dict)
    notes: dict[str, list[str]] = field(default_factory=dict)

    def render_markdown(self) -> str:
        lines = [
            "# Kết quả nhập câu hỏi",
            "",
            f"Số dòng đã điền: {self.n_rows} · nhận: {len(self.items)} · trả lại: {len(self.rejects)}",
            "",
        ]
        if self.rejects:
            lines += ["## Câu trả lại (cần sửa)", ""]
            for rid, msgs in self.rejects.items():
                lines += [f"- **{rid}**: " + "; ".join(msgs)]
            lines.append("")
        if self.notes:
            lines += ["## Lưu ý (đã nhận, nên xem lại)", ""]
            for rid, msgs in self.notes.items():
                lines += [f"- **{rid}**: " + "; ".join(msgs)]
            lines.append("")
        return "\n".join(lines)


class Reject(Exception):
    pass


def read_rows(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)[SHEET]
    title_key = {title: k for k, title, _ in COLUMNS}
    header = {j: title_key.get(str(c.value or "").strip()) for j, c in enumerate(ws[1], start=1)}
    rows = []
    for r in ws.iter_rows(min_row=2):
        row = {
            header[j]: str(c.value).strip() if c.value is not None else ""
            for j, c in enumerate(r, start=1)
            if header.get(j)
        }
        if row.get("id") and row.get("question"):
            rows.append(row)
    return rows


def _addr_tokens(text: str) -> dict[str, set[str]]:
    out: dict[str, set[str]] = defaultdict(set)
    for kind, num in _ADDR_RE.findall(text or ""):
        out[kind.lower()].add(num.lower().rstrip("."))
    return out


def address_conflict(author: str, label: str) -> bool:
    """True when the author names a unit (Điều 3) and the label names the same kind of unit with
    another number (Điều 4). A deeper or shallower address is not a conflict."""
    a, b = _addr_tokens(author), _addr_tokens(label)
    return any(kind in b and not (nums & b[kind]) for kind, nums in a.items())


class Importer:
    def __init__(
        self,
        settings: Settings,
        strata: Strata,
        resolver: Resolver,
        facts: FactIndex | None,
        author: str,
    ) -> None:
        self.settings = settings
        self.strata = strata.by_code()
        self.resolver = resolver
        self.facts = facts
        self.author = author
        self.raw = RawSource(settings)
        self.derived = Derived(settings)
        self.docs = _docs(self.raw)
        self.doc_domain = {d.document_id: d.domain for d in self.raw.registry.documents}
        self.gen = Generator(settings, settings.data_dir / "gold" / "templates_v2.yaml", 0)
        self._cells: dict[str, tuple[RawCell, list[RawCell]]] | None = None

    # ------------------------------------------------------------ evidence

    def _document(self, code: str) -> str:
        if code in self.docs:
            return self.docs[code]
        if code in self.docs.values():
            return code
        raise Reject(f"không nhận ra văn bản {code!r}; chọn trong danh sách của cột Văn bản")

    def _registry(self, doc: str, answer: str, quote: str) -> tuple[SourceRef, dict[str, Any], str]:
        wanted = dates_in(answer)
        for c in self.derived.chunks(doc):
            if c.get("chunk_type") != "registry_fact":
                continue
            text = f"{c.get('text', '')} {c.get('citation_label', '')}"
            hit = next((d for d in sorted(wanted) if d in text), None)
            if hit:
                ref = SourceRef(
                    document_id=doc,
                    file_sha256=self.raw.file(doc).sha256,
                    kind="registry",
                    line_start=0,
                    line_end=0,
                )
                q = quote if quote and quote_matches(quote, text) else hit
                return ref, c, q
        raise Reject("không có mục sổ đăng ký nào của văn bản này mang ngày ghi trong đáp án")

    def locate(
        self, code: str, address: str, quote: str, answer: str
    ) -> tuple[SourceRef, dict[str, Any], str]:
        """(source_ref, grade-2 chunk, quote) for one basis given by the author."""
        doc = self._document(code)
        if normalize_raw(address).startswith("sổ đăng ký"):
            return self._registry(doc, answer, quote)
        if not address:
            raise Reject("thiếu địa chỉ pháp lý")
        if len(normalize_raw(quote)) < 12:
            raise Reject("đoạn trích nguyên văn quá ngắn hoặc để trống")
        f = self.raw.file(doc)
        lines = self.raw.lines(doc)
        # the quote may be verbatim or verbatim fragments around "..." (quote_matches); the
        # window grows up to twenty lines so fragments of a list (a preamble, a clause that
        # enumerates replaced standards) can be found; the chunk mapping still has to cover it
        spans: list[tuple[int, int]] = [(i, i) for i in self.raw.find(doc, quote)]
        if not spans:
            for width in range(2, 21):
                for i in range(1, len(lines) - width + 2):
                    if quote_matches(quote, " ".join(lines[i - 1 : i - 1 + width])):
                        spans.append((i, i + width - 1))
                if spans:
                    break
        found: list[tuple[SourceRef, dict[str, Any]]] = []
        for ls, le in spans:
            ref = SourceRef(document_id=doc, file_sha256=f.sha256, line_start=ls, line_end=le)
            chunks = self.derived.map_to_chunks(ref)
            if chunks:
                found.append((ref, chunks[0]))
        if not found:  # text of a PDF page may be split differently from the chunk text
            for c in self.derived.chunks(doc):
                sp = c.get("span") or {}
                if quote_matches(quote, c.get("text", "")) and sp.get("line_start"):
                    ref = SourceRef(
                        document_id=doc,
                        file_sha256=f.sha256,
                        line_start=sp["line_start"],
                        line_end=sp["line_end"],
                    )
                    if self.raw.quote_in(ref, quote):
                        found.append((ref, c))
            found.sort(key=lambda x: x[0].line_end - x[0].line_start)
        if not found:
            near = nearest_line(lines, quote)
            hint = (
                f"; gần nhất trong nguồn (dòng {near[0]}): «{near[1].strip()[:220]}»"
                if near
                else ""
            )
            raise Reject(
                "không tìm thấy đoạn trích trong văn bản; chép lại đúng nguyên văn một câu"
                " (xem cột Nội dung ở sheet «Gợi ý»; được phép bỏ bớt chữ bằng «...»)" + hint
            )
        agree = [x for x in found if not address_conflict(address, x[1]["citation_label"])]
        if not agree:
            raise Reject(
                f"địa chỉ {address!r} lệch với chỗ tìm thấy đoạn trích:"
                f" {found[0][1]['citation_label']!r}"
            )
        ref, chunk = agree[0]
        sp = chunk.get("span") or {}
        ref.pdf_page, ref.printed_page = sp.get("pdf_page_start"), sp.get("printed_page")
        return ref, chunk, quote

    def _relevant(self, ref: SourceRef, chunk: dict[str, Any], *, own_ref: bool) -> RelevantChunk:
        return RelevantChunk(
            chunk_id=chunk["chunk_id"],
            citation_label=chunk["citation_label"],
            grade=2,
            source_ref=ref if own_ref else None,
        )

    # ------------------------------------------------------------ rows

    def _cell(self, key: str) -> tuple[RawCell, list[RawCell]]:
        if self._cells is None:
            self._cells = {}
            for t in B1_TABLES:
                cells = self.gen.reader.cells(t)
                for c in cells:
                    self._cells[c.cell_key] = (c, cells)
        if not key:
            return None  # type: ignore[return-value]  # only loads the cells
        if key not in self._cells:
            raise Reject(f"không tìm thấy ô {key!r} trong bảng nguồn (cột ẩn «cell» đã bị sửa?)")
        return self._cells[key]

    def _cell_by_address(self, row: dict[str, str]) -> str | None:
        """Cell key of a row outside B1 whose basis is one table cell (a trap asked about a limit):
        the quote is the printed cell and the address is the label of its rule."""
        quote = normalize_raw(row.get("quote", "")).replace(" ", "")
        address = normalize_raw(re.sub(r",?\s*tr\.\s*\d+\s*$", "", row.get("address", "")))
        if not quote or len(quote) >= 12 or "bảng" not in address or not row.get("document"):
            return None
        doc = self._document(row["document"])
        if self._cells is None:
            self._cell("")
        hits = []
        for key, (cell, _) in (self._cells or {}).items():
            if cell.document_id != doc or normalize_raw(cell.display).replace(" ", "") != quote:
                continue
            rule = self.gen._rule_for(cell)
            if rule and normalize_raw(rule["citation_label"]).endswith(address):
                hits.append(key)
        return hits[0] if len(hits) == 1 else None

    def _common(self, row: dict[str, str]) -> dict[str, Any]:
        stratum = self.strata.get(row.get("stratum", ""))
        if stratum is None:
            raise Reject("cột ẩn «stratum» trống hoặc sai; không xóa các cột ẩn")
        if row.get("split") not in ("dev", "test"):
            raise Reject("cột Tập phải là dev hoặc test")
        return {
            "schema_version": 3,
            "id": row["id"],
            "tier": "B",
            "group": stratum.group,
            "stratum": stratum.code,
            "split": row["split"],
            "question": row["question"],
            "as_of_date": AS_OF,
            "provenance": Provenance(origin="human"),
            "check_level": "self_verified",
        }

    def _verification(self, doc: str | None, quote: str | None) -> Verification:
        return Verification(
            verified_by=self.author,
            verified_at=dt.date.today(),
            method="source_file",
            source_document_id=doc,
            source_quote=quote,
        )

    def _table_item(self, row: dict[str, str], cell_key: str) -> GoldItem:
        cell, siblings = self._cell(cell_key)
        base = self.gen.item_for(cell, siblings, 1)
        if self.gen.findings:
            self.gen.findings.clear()
            raise Reject("ô này lệch giữa bảng thô và dữ liệu đã nạp; báo nhóm kỹ thuật")
        upd = self._common(row) | {
            "source": "authored_from_table",
            "tags": ["v3_round1"],
            "notes": "",
            "verification": self._verification(cell.document_id, cell.display),
        }
        stratum = self.strata[row["stratum"]]
        if stratum.group != "B1":
            if BEHAVIOURS.get(row.get("behaviour", "").lower()) != "answer":
                raise Reject("căn cứ là một ô bảng thì Hành vi mong đợi phải là «trả lời»")
            if not row.get("answer"):
                raise Reject("thiếu Đáp án ngắn")
            trap = row.get("trap") or STRATUM_TRAP.get(stratum.code)
            must_not = list(base.expected_behaviour.must_not)
            if trap in TRAPS and TRAPS[trap][1] not in must_not:
                must_not.append(TRAPS[trap][1])
            upd |= {
                "intent": INTENT.get(stratum.code, "other"),
                "expected_answer": row["answer"],
                "expected_behaviour": ExpectedBehaviour(mode="answer", must_not=must_not),
                "trap_type": trap,
                "judge_notes": row.get("note") or None,
            }
        return base.model_copy(update=upd)

    def _numbers(self, answer: str, quote: str, chunk: dict[str, Any]) -> list[ExpectedNumber]:
        text = f"{chunk.get('text', '')}\n{chunk.get('context_text', '')}"
        out = [
            ExpectedNumber(value=d, display=d, role="date", must_appear=False)
            for d in sorted(dates_in(answer) & (dates_in(text) | numbers(text)))
        ]
        plain = _values(answer) & _values(quote)
        out += [
            ExpectedNumber(value=n, display=n, role="other", must_appear=True)
            for n in sorted(plain & numbers(text))
        ]
        return out

    def _domain(self, row: dict[str, str], stratum: Stratum, docs: list[str]) -> str:
        if row.get("domain"):
            if row["domain"] not in DOMAINS:
                raise Reject(f"lĩnh vực {row['domain']!r} không hợp lệ")
            return DOMAINS[row["domain"]]
        if stratum.domain:
            return stratum.domain
        found = {self.doc_domain.get(d) for d in docs}
        if len(docs) > 1 and len(found) > 1:
            return "cross_domain"
        if found and found <= {"wastewater", "emission"} and len(found) == 1:
            return found.pop()
        raise Reject(
            "chưa chọn Lĩnh vực (bắt buộc khi căn cứ là Thông tư 10/2021 hoặc không có căn cứ)"
        )

    def _prose_item(self, row: dict[str, str]) -> GoldItem:
        stratum = self.strata[row["stratum"]]
        mode = BEHAVIOURS.get(row.get("behaviour", "").lower())
        if mode is None:
            raise Reject("cột Hành vi mong đợi phải là: trả lời / từ chối / chỉ dẫn / hỏi lại")
        if not row.get("answer"):
            raise Reject("thiếu Đáp án ngắn")
        bases: list[tuple[SourceRef, dict[str, Any], str]] = []
        for dk, ak, qk in (("document", "address", "quote"), ("document2", "address2", "quote2")):
            if not row.get(dk):
                continue
            try:
                bases.append(self.locate(row[dk], row.get(ak, ""), row.get(qk, ""), row["answer"]))
            except Reject as e:
                if mode == "clarify":  # the kit lets a «hỏi lại» row leave the basis empty
                    raise Reject(
                        f"{e} — hoặc để trống Văn bản, Địa chỉ, Trích (nhóm «hỏi lại»)"
                    ) from None
                raise
        if mode in ("answer", "pointer") and not bases:
            raise Reject(
                "câu trả lời được thì phải có Văn bản, Địa chỉ pháp lý và Trích nguyên văn"
            )
        if stratum.code == "XDOC" and len({b[0].document_id for b in bases}) < 2:
            raise Reject(
                "nhóm B4 cần căn cứ ở hai văn bản khác nhau (điền Văn bản 2, Địa chỉ 2, Trích 2)"
            )
        if mode == "abstain":
            bases = []  # a refusal cites nothing
        trap = row.get("trap") or STRATUM_TRAP.get(stratum.code)
        must_not = [TRAPS[trap][1]] if trap in TRAPS else []
        if stratum.code == "FREQ-M01":
            must_not.append("invent_frequency")
        topics = [t.strip() for t in re.split(r"[;\n]", row.get("note", "")) if t.strip()]
        if mode == "clarify" and not topics:
            raise Reject("nhóm «hỏi lại» cần ghi dữ kiện còn thiếu ở cột Ghi chú")
        domain = self._domain(row, stratum, [b[0].document_id for b in bases])
        relevant = [self._relevant(ref, c, own_ref=i > 0) for i, (ref, c, _) in enumerate(bases)]
        if mode == "clarify":
            for rc in relevant:
                rc.grade = 1
        if domain == "out_of_scope":
            relevant = []
        warnings: list[str] = []
        for _, c, _ in bases:
            if c.get("text_status") == "draft":
                warnings.append("draft_document")
            declared = f"{c.get('text') or ''} {c.get('citation_label') or ''}"
            if c.get("chunk_type") == "registry_fact" and "khai báo" in declared:
                warnings.append("admin_declared_fact")
        if domain == "out_of_scope":
            warnings.append("out_of_scope")
        if stratum.code in ("NEG-MISSING", "FREQ-M01"):
            warnings.append("missing_information")
        nums: list[ExpectedNumber] = []
        if mode == "answer":
            for _, c, q in bases:
                nums += self._numbers(row["answer"], q, c)
        # two bases in two documents; two clauses of one document are not cross-document
        factors = ["cross_document"] if len({b[0].document_id for b in bases}) > 1 else []
        if trap in ("false_premise", "injection"):
            factors.append(trap)
        answer_type = "text"
        if mode == "abstain":
            answer_type = "none"
        elif mode == "pointer":
            answer_type = "pointer"
        elif any(n.role == "date" for n in nums):
            answer_type = "date"
        elif nums:
            answer_type = "value"
        first = bases[0] if bases and relevant else None
        common = self._common(row)
        if first is None:
            common["check_level"] = "authored"  # nothing to quote for a refusal
        return GoldItem(
            **common,
            domain=domain,
            intent=INTENT.get(stratum.code, "other"),
            difficulty_factors=factors,
            source="authored_from_clause",
            answer_type=answer_type,
            expected_answer=row["answer"],
            expected_numbers=list({n.value: n for n in nums}.values()),
            source_ref=first[0] if first else None,
            relevant=relevant,
            expected_citation_labels=[rc.citation_label for rc in relevant if rc.grade == 2],
            must_abstain=mode != "answer",
            expected_behaviour=ExpectedBehaviour(
                mode=mode,
                must_not=must_not,
                pointer_labels=[rc.citation_label for rc in relevant] if mode == "pointer" else [],
                clarification_topics=topics if mode == "clarify" else [],
            ),
            expected_warnings=list(dict.fromkeys(warnings)),
            trap_type=trap,
            tags=["v3_round1"],
            judge_notes=row.get("note") or None,
            verification=self._verification(
                first[0].document_id if first else None, first[2] if first else None
            ),
        )

    def build(self, row: dict[str, str]) -> GoldItem:
        cell_key = row.get("cell") or self._cell_by_address(row)
        item = self._table_item(row, cell_key) if cell_key else self._prose_item(row)
        if self.facts is not None:
            item.fact_ids, item.legal_address = self.facts.for_item(item)
        if not item.legal_address:
            item.legal_address = [rc.citation_label for rc in item.grade2()]
        return item


def import_workbook(
    path: Path,
    settings: Settings,
    strata: Strata,
    resolver: Resolver,
    facts: FactIndex | None,
    *,
    author: str,
    existing: list[GoldItem],
) -> ImportResult:
    rows = read_rows(path)
    res = ImportResult(n_rows=len(rows))
    imp = Importer(settings, strata, resolver, facts, author)
    built: list[GoldItem] = []
    for row in rows:
        try:
            built.append(imp.build(row))
        except Reject as e:
            res.rejects[row["id"]] = [str(e)]
        except ValueError as e:  # schema refusal (id pattern, unknown value …)
            res.rejects[row["id"]] = [f"không hợp lệ: {str(e)[:300]}"]
    new_ids = {it.id for it in built}
    others = [it for it in existing if it.id not in new_ids]
    val = GoldValidator(resolver, imp.raw, None)
    for it in built:
        val.check_item(it)
    val.check_dataset(others + built)
    for f in val.rep.findings:
        if f.item_id not in new_ids:
            continue
        target = res.rejects if f.level == "E" else res.notes
        target.setdefault(f.item_id, []).append(f"{f.level}{f.code[1:]} {f.message}")
    res.items = [it for it in built if it.id not in res.rejects]
    return res
