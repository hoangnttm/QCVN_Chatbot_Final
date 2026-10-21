"""openpyxl helpers shared by the expert review pack (`scripts/review_pack.py`) and the gold-set
review workbook (`app/evaluation/gold_excel.py`): header styling, yellow reviewer columns with a
dropdown and conditional colouring, sheet finishing. Requires the optional `review` extra."""

from __future__ import annotations

from typing import Any

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

RESULT_OPTIONS = '"Đúng,Sai,Cần làm rõ"'
REVIEW_COLS = ["KẾT QUẢ", "Sửa thành / giá trị đúng", "Ghi chú chuyên gia"]
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEADER = PatternFill("solid", fgColor="D9E1F2")
RED = PatternFill("solid", fgColor="F8CBAD")
GREEN = PatternFill("solid", fgColor="C6E0B4")
GREY = PatternFill("solid", fgColor="EDEDED")


def style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True)
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def add_review_columns(
    ws: Worksheet,
    header_row: int,
    first: int,
    last: int,
    ncols: int,
    n_review: int = len(REVIEW_COLS),
) -> None:
    """Yellow reviewer columns with a dropdown + conditional colouring on the first of them."""
    col = ncols - n_review + 1
    letter = get_column_letter(col)
    dv = DataValidation(type="list", formula1=RESULT_OPTIONS, allow_blank=True)
    ws.add_data_validation(dv)
    if last >= first:
        dv.add(f"{letter}{first}:{letter}{last}")
        rng = f"{letter}{first}:{letter}{last}"
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=['"Sai"'], fill=RED)
        )
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=['"Đúng"'], fill=GREEN)
        )
    for r in range(first, last + 1):
        for c in range(col, ncols + 1):
            ws.cell(row=r, column=c).fill = YELLOW
    for c in range(col, ncols + 1):
        ws.cell(row=header_row, column=c).fill = YELLOW
        ws.cell(row=header_row, column=c).font = Font(bold=True)


def finish_sheet(ws: Worksheet, widths: list[int], header_row: int) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)


def write_table_sheet(
    ws: Worksheet,
    title_lines: list[str],
    headers: list[str],
    rows: list[list[Any]],
    widths: list[int],
    review_cols: list[str] | None = None,
    meta: dict[str, tuple[str, int, int]] | None = None,
) -> tuple[int, int]:
    """Title lines, a header row, data rows and yellow review columns; returns (first, last)."""
    review_cols = REVIEW_COLS if review_cols is None else review_cols
    for line in title_lines:
        ws.append([line])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    ws.append([])
    ws.append(headers + review_cols)
    header_row = ws.max_row
    ncols = len(headers) + len(review_cols)
    style_header(ws, header_row, ncols)
    first = header_row + 1
    for r in rows:
        ws.append(r + [""] * len(review_cols))
    last = ws.max_row
    add_review_columns(ws, header_row, first, last, ncols, n_review=len(review_cols))
    finish_sheet(ws, widths + [14, 28, 36][: len(review_cols)], header_row)
    if meta is not None:
        meta[ws.title] = (get_column_letter(ncols - len(review_cols) + 1), first, last)
    return first, last
