"""Gold v3, round 1: strata quotas, the migrated v3 files, the authoring workbook and its import.
Needs `data/derived/*` (produced by `ingest.py all`); skipped when absent."""

from __future__ import annotations

from collections import Counter
from pathlib import Path

import pytest

from app.core.config import Settings
from app.evaluation.gold_authoring import (
    COLUMNS,
    Importer,
    Reject,
    _values,
    address_conflict,
    build_workbook,
    read_rows,
)
from app.evaluation.gold_resolver import DerivedResolver
from app.evaluation.gold_schema import load_gold, load_strata
from app.evaluation.gold_source import RawSource
from app.evaluation.gold_validate import validate

REPO = Path(__file__).resolve().parents[2]
GOLD = REPO / "data" / "gold"
DERIVED = REPO / "data" / "derived" / "QCVN_40_2025" / "chunks.json"
pytestmark = pytest.mark.skipif(
    not DERIVED.exists(), reason="data/derived missing (run ingest.py all)"
)

QUOTE_TT06 = (
    "Kể từ ngày 01 tháng 01 năm 2032, các trường hợp quy định tại khoản 1 Điều 3 Thông tư này"
    " phải đáp ứng yêu cầu quy định tại QCVN 40:2025/BTNMT"
)


@pytest.fixture(scope="module")
def settings() -> Settings:
    return Settings(_env_file=None)


@pytest.fixture(scope="module")
def strata():
    return load_strata(GOLD / "strata_v3.yaml")


@pytest.fixture(scope="module")
def importer(settings, strata) -> Importer:
    return Importer(settings, strata, DerivedResolver(settings), None, "tester")


def _row(**kw: str) -> dict[str, str]:
    return {"id": "G-B3-001", "split": "dev", "stratum": "TRANSITION", "behaviour": "trả lời"} | kw


def test_round1_quotas(strata):
    hand = [s for s in strata.strata if s.tier == "B"]
    # migrated from v2 + new + 7 over-quota test rows moved to dev (D69 a, 2026-10-20)
    assert sum(s.quota_dev for s in hand) == 13 + 20 + 7
    assert sum(s.quota_test for s in hand) == 60
    by_group = Counter()
    for s in hand:
        by_group[s.group] += s.quota_test
    # B3 10 / B4 4 since 2026-10-20: two B4 rows were single-document transition questions
    assert by_group == {"B1": 15, "B2": 15, "B3": 10, "B4": 4, "B5": 6, "B6": 10}


def test_v3_files_carry_the_old_items(settings, strata):
    items = [
        it for f in ("tables", "qa", "multihop") for it in load_gold(GOLD / f"gold_{f}_v3.jsonl")
    ]
    old = [it for it in items if "v3_round1" not in it.tags]
    old = [it for it in old if it.split != "test"]
    assert Counter(it.split for it in old) == {"dev": 116, "val": 291}  # 404 - 113 rows of D53
    assert all(it.schema_version == 3 and it.group for it in old)
    answerable = [it for it in old if it.expected_behaviour.mode == "answer"]
    assert all(it.legal_address for it in answerable)
    rep = validate(items, DerivedResolver(settings), RawSource(settings), strata)
    assert not rep.errors


def test_address_conflict():
    label = "Thông tư 06/2025/TT-BTNMT, Điều 4, Khoản 2"
    assert not address_conflict("khoản 2 Điều 4", label)
    assert not address_conflict("Điều 4", label)  # shallower is fine
    assert not address_conflict("Điều 4, Khoản 2, Điểm a", label)  # deeper is fine
    assert address_conflict("Điều 5, Khoản 2", label)
    assert not address_conflict("Mục 1.3.3.1", "QCVN 40:2025/BTNMT, Mục 1.3.3.1")


def test_values_ignore_references():
    text = (
        "Cột B của Bảng 1 và Bảng 2 theo QCVN 40:2025/BTNMT, Thông tư 06/2025/TT-BTNMT; 0,05 mg/L"
    )
    assert _values(text) == {"0.05"}


def test_import_prose_row(importer, settings, strata):
    item = importer.build(
        _row(
            question="Cơ sở đang hoạt động phải đáp ứng quy chuẩn nước thải mới từ khi nào?",
            answer="Từ ngày 01/01/2032",
            document="06/2025/TT-BTNMT",
            address="Điều 4, Khoản 2",
            quote=QUOTE_TT06,
        )
    )
    assert item.group == "B3" and item.domain == "wastewater" and item.intent == "applicability"
    assert [rc.chunk_id for rc in item.grade2()] == ["chunk:TT_06_2025/Article4/Clause2"]
    assert item.legal_address == ["Thông tư 06/2025/TT-BTNMT, Điều 4, Khoản 2"]
    assert [n.value for n in item.expected_numbers] == ["2032-01-01"]
    rep = validate([item], DerivedResolver(settings), RawSource(settings), None)
    assert not rep.errors


def test_import_hands_back_bad_rows(importer):
    with pytest.raises(Reject, match="lệch"):
        importer.build(
            _row(
                question="q?",
                answer="01/01/2032",
                document="06/2025/TT-BTNMT",
                address="Điều 5, Khoản 2",
                quote=QUOTE_TT06,
            )
        )
    with pytest.raises(Reject, match="không tìm thấy đoạn trích"):
        importer.build(
            _row(
                question="q?",
                answer="x",
                document="06/2025/TT-BTNMT",
                address="Điều 4",
                quote="một câu không hề có trong thông tư này",
            )
        )
    with pytest.raises(Reject, match="Ghi chú"):
        importer.build(
            _row(
                stratum="CLARIFY",
                question="COD bao nhiêu?",
                answer="hỏi lại",
                behaviour="hỏi lại",
                domain="nước thải",
            )
        )


def test_import_refusal_and_clarify(importer):
    oos = importer.build(
        _row(
            id="G-B6-001",
            stratum="NEG-OOS",
            question="Mức ồn cho phép là bao nhiêu?",
            answer="Ngoài phạm vi",
            behaviour="từ chối",
            note="tiếng ồn",
        )
    )
    assert oos.must_abstain and oos.domain == "out_of_scope" and not oos.relevant
    assert oos.trap_type == "out_of_scope" and oos.check_level == "authored"
    ask = importer.build(
        _row(
            id="G-B5-001",
            stratum="CLARIFY",
            question="COD nước thải tối đa bao nhiêu?",
            answer="hỏi lại",
            behaviour="hỏi lại",
            domain="nước thải",
            note="lưu lượng xả thải; cột",
        )
    )
    assert ask.expected_behaviour.mode == "clarify"
    assert ask.expected_behaviour.clarification_topics == ["lưu lượng xả thải", "cột"]


def test_import_degree_sign_and_trap_on_a_table_cell(importer):
    nm3 = importer.build(  # the source prints 25°C in LaTeX, the author typed "25 oC"
        _row(
            id="G-B2-007",
            stratum="DEF-SCOPE",
            question="Nm³ trong quy chuẩn khí thải lấy nhiệt độ chuẩn bao nhiêu?",
            answer="25 °C và áp suất tuyệt đối 760 mm Thủy ngân.",
            document="QCVN 19:2024/BTNMT",
            address="Mục 1.3.10, tr. 4",
            quote="Mét khối khí thải chuẩn (Nm3) là mét khối khí thải ở nhiệt độ 25 oC",
        )
    )
    assert nm3.legal_address == ["QCVN 19:2024/BTNMT, Mục 1.3.10"]
    trap = importer.build(  # a trap whose basis is one cell: found by its address and printed value
        _row(
            id="G-B6-010",
            stratum="NEG-TRAP",
            trap="pass_fail_request",
            question="Lò hơi dùng khí sinh học, cột A, bụi đo được 18 mg/Nm³. Nhà máy đạt chưa?",
            answer="Ngưỡng bụi: ≤ 20 mg/Nm³, ôxy tham chiếu 4 %. Không kết luận đạt hay không đạt.",
            document="QCVN 19:2024/BTNMT",
            address="Mục 2.2, Bảng 2, TT 1 (Bụi), Thiết bị 1.5, Cột A",
            quote="≤ 20(4)",
        )
    )
    assert trap.group == "B6" and trap.trap_type == "pass_fail_request"
    assert trap.target_rule_codes and [n.value for n in trap.expected_numbers] == ["20", "4"]
    assert "pass_fail_verdict" in trap.expected_behaviour.must_not


def test_workbook_round_trip(settings, strata, importer, tmp_path):
    from openpyxl import load_workbook

    existing = [  # the set as it was when the kit was issued, before the hand-written items
        it
        for f in ("tables", "qa", "multihop")
        for it in load_gold(GOLD / f"gold_{f}_v3.jsonl")
        if "v3_round1" not in it.tags
    ]
    out = tmp_path / "kit.xlsx"
    # 20 dev + 60 test rows by D46; D69 a (2026-10-20) raised the dev quotas by 7 -> 27 + 60
    assert build_workbook(settings, strata, existing, out, seed=None) == 87
    wb = load_workbook(out)
    ws = wb["Soạn"]
    keys = [k for k, _, _ in COLUMNS]
    rows = [dict(zip(keys, r, strict=True)) for r in ws.iter_rows(min_row=2, values_only=True)]
    assert Counter(r["split"] for r in rows) == {"dev": 27, "test": 60}
    b1 = [r for r in rows if r["group"] == "B1"]
    assert len(b1) == 20 and len({r["cell"] for r in b1}) == 20
    dev_rows = {it.cluster_id for it in existing if it.split == "dev"}
    for r in b1:  # no test cell from a row whose template questions are in dev
        if r["split"] == "test":
            assert (
                r["cell"].rsplit(".", 2 if "F" in r["cell"].split(".")[-2] else 1)[0]
                not in dev_rows
            )
    assert read_rows(out) == []  # nothing filled in yet
    ws.cell(
        2,
        keys.index("question") + 1,
        "Nhà máy xả 3.500 m3/ngày đêm, cột A thì COD tối đa bao nhiêu?",
    )
    wb.save(out)
    (row,) = read_rows(out)
    item = importer.build(row)
    assert item.id == "G-B1-001" and item.tier == "B" and item.target_rule_codes
    assert item.forbidden_numbers and item.source == "authored_from_table"


def test_d53_reserved_rows_and_template_test(settings):
    import yaml

    from app.evaluation.gold_v3_test import T1_QUOTA, T2_QUOTA, generate_test, load_reserve

    reserve = load_reserve(GOLD / "test_rows_v3.yaml")
    existing = [
        it
        for f in ("tables", "multihop")
        for it in load_gold(GOLD / f"gold_{f}_v3.jsonl")
        if it.split != "test"
    ]
    reserved = set(reserve["reserved_rows"])
    assert not reserved & {it.cluster_id for it in existing}  # gone from dev and val
    assert set(reserve["hand_rows"]["test"]) & reserved
    spec = yaml.safe_load((GOLD / "multihop_v2.yaml").read_text(encoding="utf-8"))
    new = {
        f"T1:{k}": ["Xin hỏi: " + t["text"] for t in v[:2]]
        for k, v in yaml.safe_load((GOLD / "templates_v2.yaml").read_text(encoding="utf-8"))[
            "templates"
        ].items()
    } | {
        f"T2:{p}:{d}": ["Xin hỏi: " + t["text"] for t in tpls[:1]]
        for p, per in spec["templates"].items()
        for d, tpls in per.items()
    }
    t1, t2, notes = generate_test(
        settings, GOLD / "multihop_v2.yaml", new, reserve, existing, None, seed=20261006
    )
    assert not notes
    assert len(t1) == sum(T1_QUOTA.values()) and len(t2) == sum(
        sum(q.values()) for q in T2_QUOTA.values()
    )
    assert Counter(it.domain for it in t1) == T1_QUOTA
    used = {it.cluster_id for it in existing} | set(reserve["hand_rows"]["dev"])
    hand = set(reserve["hand_cells"])
    for it in t1 + t2:
        assert it.split == "test" and it.cluster_id not in used
        assert not set(it.target_rule_codes) & hand
        assert it.question.startswith("Xin hỏi")
        if it.domain == "wastewater":
            assert it.cluster_id in reserved
    rep = validate(t1 + t2, DerivedResolver(settings), RawSource(settings), None)
    assert not rep.errors


def test_frame_labels_and_score(settings):
    from app.evaluation.gold_frame import (
        PredictedFrame,
        document_keys,
        score_item,
        summarize,
        template_frame,
    )

    assert document_keys("Thông tư 06/2025 và QCVN 40:2025/BTNMT, qcvn 19:2024") == [
        "TT6/2025",
        "QCVN40:2025",
        "QCVN19:2024",
    ]
    assert document_keys("theo 45/2024/TT-BTNMT") == ["TT45/2024"]
    resolver = DerivedResolver(settings)
    by_tag = {}
    for it in load_gold(GOLD / "gold_multihop_v3.jsonl"):
        by_tag.setdefault(next(t for t in it.tags if t.startswith("pattern_")), it)
    frames = {tag: template_frame(it, resolver) for tag, it in by_tag.items()}
    dz, tr, rp = frames["pattern_dz"], frames["pattern_tr"], frames["pattern_rp"]
    # the default column of an unzoned facility is what the answer needs, not what is asked
    assert dz.intents == ["threshold", "applicability"]
    assert "zone" not in {s.dimension for s in dz.slots}
    assert ("date", "2032-01-01") in {(s.dimension, s.value) for s in tr.slots}
    assert rp.intents == ["threshold", "effectivity"]
    assert ("document", "QCVN40:2011") in {(s.dimension, s.value) for s in rp.slots}

    item = by_tag["pattern_tr"].model_copy(update={"frame": tr})
    gold = {(s.dimension, s.value, s.negated) for s in tr.slots}
    right = score_item(item, PredictedFrame(item.domain, ["threshold", "applicability"], gold))
    assert right["both"] == 1.0 and right["intent_f1"] == 1.0
    assert all(fp == 0 and fn == 0 for _, fp, fn in right["slots"].values())
    # main intent decides the gate (D55); a missing secondary intent only lowers the set F1
    main_only = score_item(item, PredictedFrame(item.domain, ["threshold"], set()))
    assert main_only["both"] == 1.0 and main_only["intent_f1"] < 1.0
    wrong = score_item(item, PredictedFrame("unknown", ["applicability"], set()))
    assert wrong["both"] == 0.0
    assert summarize([right, main_only, wrong])["all"]["both"]["mean"] == pytest.approx(2 / 3)


def test_every_dev_and_val_item_has_a_frame():
    items = [
        it for f in ("tables", "qa", "multihop") for it in load_gold(GOLD / f"gold_{f}_v3.jsonl")
    ]
    assert all(it.frame for it in items if it.split in ("dev", "val"))
    hand = [it for it in items if it.tier == "B"]
    assert all(it.frame.intents_by == "reviewed" for it in hand)
