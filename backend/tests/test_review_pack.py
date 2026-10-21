"""Expert review loop: pack generation is deterministic and risk-first; a filled workbook flows
back into review_status + review_decisions, including the §11 bulk approval."""

import importlib.util
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.db import models as m
from app.db.loader import load_document
from app.db.session import get_sessionmaker
from app.registry.loader import load_registry

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"


def _load(name: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="module")
def pack(db_settings, tmp_path_factory):
    reg = load_registry(settings=db_settings)
    maker = get_sessionmaker(db_settings)
    with maker.begin() as session:
        for doc_id in ("TT_06_2025", "QCVN_40_2025"):
            load_document(doc_id, reg, db_settings, session)
    review_pack = _load("review_pack")
    out = tmp_path_factory.mktemp("review") / "pack.xlsx"
    review_pack.SHEET_META.clear()
    review_pack.build(out, 10, db_settings)
    return out


def test_pack_sampling_is_risk_first_and_stable(pack, db_settings):
    wb = load_workbook(pack)
    assert {"0-Hướng dẫn", "1-Tổng hợp", "2b-QCVN40 B2", "6-Quan hệ", "7-Chuyển tiếp"} <= set(
        wb.sheetnames
    )
    ws = wb["2b-QCVN40 B2"]
    header = [c.value for c in ws[4]]
    assert header[:2] == ["Mã dòng", "Nhãn trích dẫn"] and header[-3] == "KẾT QUẢ"
    ids = [ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1)]
    assert len(ids) == 10
    # risk-first: the receiving-body / sector conditioned T-P rows and the pH range are in the sample
    assert {"QCVN40_2025.B2.r4.1", "QCVN40_2025.B2.r4.2", "QCVN40_2025.B2.r1"} <= set(ids)
    ph = ws.cell(row=ids.index("QCVN40_2025.B2.r1") + 5, column=header.index("Giá trị A") + 1).value
    assert ph == "6.0 – 9.0"
    # deterministic
    review_pack = sys.modules["review_pack"]
    review_pack.SHEET_META.clear()
    again = pack.parent / "pack2.xlsx"
    review_pack.build(again, 10, db_settings)
    ws2 = load_workbook(again)["2b-QCVN40 B2"]
    assert [ws2.cell(row=r, column=1).value for r in range(5, ws2.max_row + 1)] == ids
    summary = wb["1-Tổng hợp"]
    rows = {
        summary.cell(row=r, column=1).value: summary.cell(row=r, column=5).value
        for r in range(2, 8)
    }
    assert rows["2a-QCVN40 B1"] == 24 and rows["2b-QCVN40 B2"] == 189


def test_import_updates_status_and_logs_decisions(pack, db_settings):
    wb = load_workbook(pack)
    ws = wb["2a-QCVN40 B1"]
    res_col = [c.value for c in ws[4]].index("KẾT QUẢ") + 1
    for r in range(5, ws.max_row + 1):  # all 4 rows correct -> bulk approval of Bảng 1
        ws.cell(row=r, column=res_col, value="Đúng")
    ws = wb["2b-QCVN40 B2"]
    res_col = [c.value for c in ws[4]].index("KẾT QUẢ") + 1
    ws.cell(row=5, column=res_col, value="Sai")
    ws.cell(row=5, column=res_col + 1, value="giá trị đúng theo bản ký")
    ws.cell(row=6, column=res_col, value="Đúng")
    ws = wb["7-Chuyển tiếp"]
    res_col = [c.value for c in ws[4]].index("KẾT QUẢ") + 1
    ws.cell(row=5, column=res_col, value="Cần làm rõ")
    filled = pack.parent / "filled.xlsx"
    wb.save(filled)

    review_import = _load("review_import")
    decisions = review_import.read_decisions(filled)
    assert sum(1 for d in decisions if d["decision"]) == 7
    summary = review_import.apply(decisions, "Chuyên gia A", filled.name, True, db_settings)
    assert summary["decisions"] == 7 and summary["bulk"] == [
        {"table": "QCVN40_2025.B1", "sample": 4, "rules": 24}
    ]
    maker = get_sessionmaker(db_settings)
    with maker() as session:
        by_status = dict(
            session.execute(
                select(m.LimitRule.review_status, func.count())
                .where(m.LimitRule.document_id == "QCVN_40_2025")
                .group_by(m.LimitRule.review_status)
            ).all()
        )
        assert by_status["approved"] == 24 + 3 and by_status["rejected"] == 3
        assert by_status["pending"] == 213 - 30
        rejected = session.execute(
            select(m.ReviewDecision).where(m.ReviewDecision.decision == "rejected")
        ).scalar_one()
        assert (
            rejected.correction == "giá trị đúng theo bản ký"
            and rejected.reviewer == "Chuyên gia A"
        )
        bulk = session.execute(
            select(m.ReviewDecision).where(m.ReviewDecision.target_type == "limit_rules:table")
        ).scalar_one()
        assert bulk.target_id == "QCVN40_2025.B1" and "4 sampled rows" in bulk.note
        recheck = session.execute(
            select(m.ReviewDecision).where(m.ReviewDecision.decision == "needs_recheck")
        ).scalar_one()
        assert recheck.target_type == "transition_rules"
