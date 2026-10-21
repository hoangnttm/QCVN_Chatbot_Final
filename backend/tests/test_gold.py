"""Gold set v2 tooling: schema, difficulty, label matching, validator codes on synthetic items,
v1 migration against the derived data, split assignment, Excel round trip. Needs
`data/derived/*` (produced by `ingest.py all`); skipped when absent."""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

import pytest

from app.core.config import Settings
from app.evaluation import gold_stats
from app.evaluation.gold_agreement import cohen_kappa, compare, wilson_ci
from app.evaluation.gold_candidates import candidate_rows
from app.evaluation.gold_migrate import migrate_v1
from app.evaluation.gold_resolver import DerivedResolver
from app.evaluation.gold_schema import (
    ExpectedBehaviour,
    ExpectedConditions,
    ExpectedNumber,
    GoldItem,
    RelevantChunk,
    SourceRef,
    Verification,
    compute_difficulty,
    dump_gold,
    load_gold,
    load_strata,
)
from app.evaluation.gold_source import Derived, RawSource, normalize_raw
from app.evaluation.gold_split import assign_splits
from app.evaluation.gold_validate import dates_in, label_matches, validate

REPO = Path(__file__).resolve().parents[2]
DERIVED = REPO / "data" / "derived" / "QCVN_40_2025" / "chunks.json"
pytestmark = pytest.mark.skipif(
    not DERIVED.exists(), reason="data/derived missing (run ingest.py all)"
)


@pytest.fixture(scope="module")
def settings() -> Settings:
    return Settings(_env_file=None)


@pytest.fixture(scope="module")
def resolver(settings) -> DerivedResolver:
    return DerivedResolver(settings)


@pytest.fixture(scope="module")
def raw(settings) -> RawSource:
    return RawSource(settings)


@pytest.fixture(scope="module")
def strata():
    return load_strata(REPO / "data" / "gold" / "strata_v2.yaml")


# ---------------------------------------------------------------- pure functions


def test_compute_difficulty():
    assert compute_difficulty([]) == "easy"
    assert compute_difficulty(["zone", "flow_band"]) == "easy"
    assert compute_difficulty(["zone", "receiving_body", "sector_scope"]) == "medium"
    assert compute_difficulty(["zone", "equipment_depth3", "fuel", "capacity", "o2_ref"]) == "hard"
    assert compute_difficulty(["zone", "flow_band", "footnote"]) == "medium"
    assert compute_difficulty(["injection", "zone"]) == "medium"  # HARD factors count twice


def test_label_matches_is_separator_aware():
    row = "QCVN 40:2025/BTNMT, Mục 2.1, Bảng 1, TT 1"
    assert label_matches(row, row)
    assert label_matches(row + ", F > 2000, Cột B", row)
    assert not label_matches("QCVN 40:2025/BTNMT, Mục 2.1, Bảng 1, TT 12, Cột B", row)
    assert not label_matches("QCVN 40:2025/BTNMT, Mục 2.1, Bảng 1", row)  # reverse direction


def test_dates_and_normalize():
    assert dates_in("có hiệu lực từ 01/9/2025 và ngày 30 tháng 12 năm 2024") == {
        "2025-09-01",
        "2024-12-30",
    }
    assert normalize_raw("| 1 | BOD | $mg/L$ | $\\le 40$ |") == "| 1 | bod | mg/l | ≤ 40 |"


def test_wilson_and_kappa():
    lo, hi = wilson_ci(9, 10)
    assert 0.55 < lo < 0.7 and 0.98 < hi <= 1.0
    assert cohen_kappa(
        [(True, True), (False, False), (True, True), (False, True)]
    ) == pytest.approx(0.5)
    assert cohen_kappa([]) is None


# ---------------------------------------------------------------- schema io


def test_schema_roundtrip(tmp_path):
    item = GoldItem(
        id="G-WW-900",
        question="pH nước thải công nghiệp cột B?",
        domain="wastewater",
        intent="threshold",
        expected_answer="6 – 9",
        expected_numbers=[
            ExpectedNumber(value="6", role="range_min"),
            ExpectedNumber(value="9", role="range_max"),
        ],
        expected_conditions=ExpectedConditions(zone="B"),
        difficulty_factors=["zone"],
        difficulty="easy",
    )
    p = tmp_path / "g.jsonl"
    dump_gold([item], p)
    back = load_gold(p)
    assert back[0] == item
    with pytest.raises(ValueError):
        GoldItem.model_validate({**item.model_dump(mode="json"), "difficulty_factors": ["bogus"]})
    with pytest.raises(ValueError):
        GoldItem.model_validate({**item.model_dump(mode="json"), "unknown_field": 1})


# ---------------------------------------------------------------- derived data


def _bod_item(resolver) -> GoldItem:
    rule = resolver.rule("QCVN40_2025.B1.r1.F>2000.B")
    chunk = resolver.chunk("chunk:QCVN40_2025.B1.r1")
    sp = chunk["span"]
    return GoldItem(
        id="G-WW-901",
        question="Giới hạn BOD5 trong nước thải công nghiệp cột B, lưu lượng 3000 m3/ngày là bao nhiêu?",
        domain="wastewater",
        intent="threshold",
        as_of_date=dt.date(2026, 1, 15),
        split="dev",
        stratum="WW-B1-FLOWZONE",
        expected_answer="≤ 50 mg/L (F > 2000, Cột B)",
        expected_numbers=[ExpectedNumber(value="50", display="50", unit="mg/L", role="limit")],
        expected_conditions=ExpectedConditions(zone="B", flow_band="F>2000"),
        target_rule_codes=[rule["rule_code"]],
        source_ref=SourceRef(
            document_id="QCVN_40_2025",
            file_sha256=sp["file_hash"],
            line_start=sp["line_start"],
            line_end=sp["line_end"],
            table_code="QCVN40_2025.B1",
            row_no="1",
            column_path=["F>2000", "B"],
        ),
        relevant=[
            RelevantChunk(
                chunk_id=chunk["chunk_id"], citation_label=chunk["citation_label"], grade=2
            )
        ],
        expected_citation_labels=[rule["citation_label"]],
        difficulty_factors=["zone", "flow_band"],
        difficulty="easy",
        verification=Verification(
            verified_by="tester",
            verified_at=dt.date(2026, 9, 25),
            method="source_file",
            source_document_id="QCVN_40_2025",
            source_quote="Nhu cầu ôxy sinh hóa",
        ),
        check_level="self_verified",
    )


def test_validator_accepts_a_correct_item(resolver, raw):
    rep = validate([_bod_item(resolver)], resolver, raw)
    assert rep.ok, [str(f) for f in rep.errors]


def test_validator_codes_on_bad_items(resolver, raw):
    good = _bod_item(resolver)
    bad_label = good.model_copy(deep=True)
    bad_label.relevant[0].citation_label = "QCVN 40:2025/BTNMT, Mục 2.1, Bảng 1, TT 12"
    bad_number = good.model_copy(deep=True)
    bad_number.expected_numbers = [ExpectedNumber(value="55", role="limit")]
    ambiguous = good.model_copy(deep=True)
    ambiguous.expected_conditions = ExpectedConditions(zone="B")  # flow band missing
    wrong_abstain = good.model_copy(deep=True)
    wrong_abstain.must_abstain = True
    missing_quote = good.model_copy(deep=True)
    missing_quote.verification = Verification(
        verified_by="t", verified_at=dt.date(2026, 9, 25), source_quote="không có trong nguồn"
    )
    bad_diff = good.model_copy(deep=True)
    bad_diff.difficulty = "hard"
    para = good.model_copy(deep=True)
    para.source = "paraphrase_llm_verified"
    codes = {}
    for name, it in [
        ("label", bad_label),
        ("number", bad_number),
        ("ambiguous", ambiguous),
        ("abstain", wrong_abstain),
        ("quote", missing_quote),
        ("diff", bad_diff),
        ("para", para),
    ]:
        rep = validate([it], resolver, raw)
        codes[name] = {f"{f.level}{f.code[1:]}" for f in rep.errors}
    assert "E06" in codes["label"]
    assert "E09" in codes["number"] or "E10" in codes["number"]
    assert "E10" in codes["ambiguous"]
    assert "E11" in codes["abstain"]
    assert "E13" in codes["quote"]
    assert "E21" in codes["diff"]
    assert "E14" in codes["para"]


def test_validator_dataset_checks(resolver, raw, strata):
    a = _bod_item(resolver)
    b = a.model_copy(deep=True)
    b.id = "G-WW-902"
    rep = validate([a, b], resolver, raw, strata)
    assert any(f.code == "E18" for f in rep.findings)  # identical normalized question
    c = a.model_copy(deep=True)
    c.id = "G-WW-903"
    c.paired_with = "G-WW-904"
    rep = validate([c], resolver, raw, strata)
    assert any(f.code == "E17" for f in rep.errors)


def test_migrate_v1_resolves_every_label(resolver, raw, strata, tmp_path):
    items, problems = migrate_v1(REPO / "data" / "gold" / "gold_qa_v1.jsonl", resolver)
    assert problems == []
    assert len(items) == 13 and all(
        it.split == "dev" and it.source == "v1_migrated" for it in items
    )
    rep = validate(items, resolver, raw, strata)
    hard = {f"{f.level}{f.code[1:]}" for f in rep.errors}
    assert not hard, [str(f) for f in rep.errors]
    eff = next(it for it in items if it.id == "G-EFF-001")
    assert eff.source_ref and eff.source_ref.kind == "registry"
    assert eff.expected_citation_labels[0].endswith(")")  # audit suffix copied byte-for-byte
    # every grade-2 chunk covers its source_ref (mapping layer consistent)
    der = Derived(Settings(_env_file=None))
    for it in items:
        for rc in it.grade2():
            ch = resolver.chunk(rc.chunk_id)
            assert der.chunk_covers(ch, rc.source_ref or it.source_ref)


def test_assign_splits_is_stable_and_respects_twins(resolver, strata):
    base = _bod_item(resolver)
    items = []
    for i in range(12):
        it = base.model_copy(deep=True)
        it.id = f"G-WW-{910 + i}"
        it.split = None
        it.question = f"{base.question} biến thể {i}"
        items.append(it)
    items[0].paired_with = items[1].id
    items[1].paired_with = items[0].id
    a = assign_splits([it.model_copy(deep=True) for it in items], strata)
    b = assign_splits([it.model_copy(deep=True) for it in items], strata)
    assert [x.split for x in a] == [x.split for x in b]
    assert sum(1 for x in a if x.split == "dev") == strata.by_code()["WW-B1-FLOWZONE"].quota_dev
    assert a[0].split == a[1].split
    again = assign_splits(a, strata)
    assert [x.split for x in again] == [x.split for x in a]  # never reassigns


def test_candidates_cover_zones_and_flow_bands(resolver, strata):
    s = strata.by_code()["WW-B1-FLOWZONE"]
    cands = candidate_rows(resolver, s, seed=strata.seed, n=8)
    assert cands and all(c.chunk_id and c.skeleton["target_rule_codes"] for c in cands)
    cells = {x for c in cands for x in c.coverage_cells}
    assert {"zone=A", "zone=B", "zone=C", "flow=F<=2000", "flow=F>2000"} <= cells
    again = candidate_rows(resolver, s, seed=strata.seed, n=8)
    assert [c.key for c in again] == [c.key for c in cands]


def test_stats_render(resolver):
    items, _ = migrate_v1(REPO / "data" / "gold" / "gold_qa_v1.jsonl", resolver)
    text = gold_stats.render_markdown(items, resolver)
    assert "Tổng: 13 câu" in text and "G-WW-001" in text
    hidden = gold_stats.render_markdown(
        [it.model_copy(update={"split": "test"}) for it in items], resolver
    )
    assert "G-WW-001" not in hidden.split("## Câu hỏi")[1]


def test_excel_roundtrip(resolver, tmp_path):
    pytest.importorskip("openpyxl")
    from app.evaluation.gold_excel import apply_review, export_review, read_review

    items, _ = migrate_v1(REPO / "data" / "gold" / "gold_qa_v1.jsonl", resolver)
    out = tmp_path / "gold.xlsx"
    export_review(items, out, mode="expert")
    decisions = read_review(out)
    assert [d.item_id for d in decisions] == [it.id for it in items]
    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb["2-Câu hỏi"]
    hrow = next(r for r in range(1, 10) if ws.cell(row=r, column=1).value == "Mã")
    header = [c.value for c in ws[hrow]]
    col = header.index("KẾT QUẢ") + 1
    ws.cell(row=hrow + 1, column=col, value="Sai")
    ws.cell(row=hrow + 1, column=col + 1, value="≤ 45 mg/L")
    wb.save(out)
    decisions = read_review(out)
    n = apply_review(items, decisions, "Chuyên gia B")
    assert (
        n == 1
        and items[0].status == "rejected"
        and items[0].verification.expert_correction == "≤ 45 mg/L"
    )
    export_review(items, tmp_path / "second.xlsx", mode="second-check")
    ws2 = load_workbook(tmp_path / "second.xlsx")["2-Câu hỏi"]
    hrow2 = next(r for r in range(1, 10) if ws2.cell(row=r, column=1).value == "Mã")
    header2 = [c.value for c in ws2[hrow2]]
    assert "Đáp án kỳ vọng" not in header2 and "Đáp án của bạn" in header2


def test_agreement_report(resolver):
    items, _ = migrate_v1(REPO / "data" / "gold" / "gold_qa_v1.jsonl", resolver)
    bod = next(it for it in items if it.id == "G-WW-001")
    bod.expected_numbers = [ExpectedNumber(value="50", role="limit")]
    seconds = [
        {
            "item_id": "G-WW-001",
            "numbers": "50",
            "must_abstain": "Không",
            "labels": "Bảng 1, TT 1, Cột B",
        },
        {"item_id": "G-OOS-001", "answer": "ngoài phạm vi", "must_abstain": "Có"},
        {"item_id": "G-WW-003", "numbers": "6; 9,5", "must_abstain": "Không"},
    ]
    rep = compare(items, seconds)
    assert rep.n == 3 and rep.abstain_agree == 3 and rep.numbers_agree >= 2
    assert rep.kappa_abstain == 1.0


def test_source_quote_check(raw, resolver):
    chunk = resolver.chunk("chunk:QCVN40_2025.B1.r1")
    sp = chunk["span"]
    ref = SourceRef(
        document_id="QCVN_40_2025",
        file_sha256=sp["file_hash"],
        line_start=sp["line_start"],
        line_end=sp["line_end"],
    )
    assert raw.sha256_ok("QCVN_40_2025")
    assert raw.quote_in(ref, "BOD5 ở 20°C")
    assert raw.quote_in(ref, "≤ 40")
    assert not raw.quote_in(ref, "Thủy ngân")
    behaviour = ExpectedBehaviour(mode="abstain")
    assert json.loads(behaviour.model_dump_json())["mode"] == "abstain"


def test_quote_matches_tolerates_ellipsis_and_punctuation():
    from app.evaluation.gold_source import loose, nearest_line, quote_fragments, quote_matches

    hay = "| 6 | Chế biến thực phẩm | C 10 | Amoni, tổng Nitơ, tổng Phốt pho. | ghi chú |"
    assert loose("C 10: Amoni, tổng Nitơ") == "c 10 amoni tổng nitơ"
    assert quote_matches("C 10: Amoni, tổng Nitơ, tổng Phốt pho.", hay)
    assert quote_fragments("Cột A quy định... khi xả ra nguồn nước") == [
        "cột a quy định",
        "khi xả ra nguồn nước",
    ]
    assert quote_matches("Chế biến thực phẩm ... tổng Phốt pho", hay)
    assert quote_matches("Chế biến thực phẩm … tổng Phốt pho", hay)
    assert not quote_matches("tổng Phốt pho ... Chế biến thực phẩm", hay)  # out of order
    assert not quote_matches("", hay)
    assert not quote_matches("...", hay)
    assert not quote_matches("Chế biến thực phẩm ... Thủy ngân", hay)
    # "≤40", "≤ 40" and "<= 40" are one reading (table cells copied without the space)
    row = r"| 2 | Nhiệt độ | $^\circ\text{C}$ | $\le 40$ | $\le 40$ | $\le 40$ |"
    assert quote_matches("Nhiệt độ ... ≤40 ... ≤40 ... ≤40", row)
    assert quote_matches("Nhiệt độ ... <= 40", row)
    assert loose("≤40") == loose("≤ 40") == "≤ 40"
    lines = [
        "ngắn",
        "Điều 1. Ban hành kèm theo Thông tư này Quy chuẩn kỹ thuật quốc gia",
        "Cơ sở đã đi vào vận hành trước ngày Thông tư này có hiệu lực thi hành",
    ]
    near = nearest_line(lines, "Cơ sở đã đi vào vận hành ... có hiệu lực")
    assert near is not None and near[0] == 3
    assert nearest_line(lines, "zzzz qqqq xxxx") is None
